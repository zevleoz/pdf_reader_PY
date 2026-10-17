#!/usr/bin/env python3.11
# -*- coding: utf-8 -*-
"""6 A.M. 算法规格文档：Markdown → DOCX（腾讯文档可导入）——一次性转换脚本。

用法（homebrew python3.11）：
    /opt/homebrew/bin/python3.11 docs/build_tencent_docx.py            # 生成（PNG 已存在则跳过渲染）
    /opt/homebrew/bin/python3.11 docs/build_tencent_docx.py --force-render  # 强制重新渲染 PNG
    /opt/homebrew/bin/python3.11 docs/build_tencent_docx.py \
        --src docs/6AM_model_algorithm_simple.md \
        --out docs/6AM_model_algorithm_simple_tencent.docx   # 简化版（复用同一批 PNG）

依赖：
    /opt/homebrew/bin/python3.11 -m pip install python-docx pillow

流程：
    1. 从 md 提取 7 个 mermaid 代码块，经 mermaid.ink 渲染为 PNG 存入 docs/_6am_diagrams/；
       每个 PNG 校验非空、可打开（Pillow）、尺寸合理；
    2. 解析 md（标题层级 / 段落 / 引用 / 有序无序列表 / 表格 / 代码块 / 行内加粗与行内代码）；
    3. 生成 docx：mermaid 块替换为对应 PNG 嵌入图 + 图注，表格用带边框表格，代码块转等宽纯文本段落；
    4. 重新打开 docx 做结构校验（图片数 / 表格数 / 标题数）并打印结果。

注意：本脚本只读 6AM_model_algorithm.md，不改动原文件。
"""

from __future__ import annotations

import argparse
import base64
import io
import re
import sys
import urllib.request
from pathlib import Path

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
    from PIL import Image
except ImportError as exc:  # pragma: no cover
    sys.exit(f"缺少依赖：{exc}\n请先执行：/opt/homebrew/bin/python3.11 -m pip install python-docx pillow")

DOCS_DIR = Path(__file__).resolve().parent
SRC_MD = DOCS_DIR / "6AM_model_algorithm.md"
OUT_DOCX = DOCS_DIR / "6AM_model_algorithm_tencent.docx"
DIAGRAM_DIR = DOCS_DIR / "_6am_diagrams"

# mermaid 块顺序 → PNG 文件名（与 §3.1、§4.1–§4.6 一一对应）
DIAGRAM_NAMES = [
    "01_mindmap_overview",
    "02_a1_autonomy",
    "03_a2_alignment",
    "04_a3_availability",
    "05_a4_assurance",
    "06_a5_accessibility",
    "07_a6_activation_support",
]

MERMAID_RE = re.compile(r"```mermaid[ \t]*\n(.*?)```", re.S)
HEADING_RE = re.compile(r"^(#{1,6})\s+(.*?)\s*$")
UL_RE = re.compile(r"^(\s*)[-*]\s+(.*)$")
OL_RE = re.compile(r"^(\s*)(\d+)\.\s+(.*)$")
QUOTE_RE = re.compile(r"^>\s?(.*)$")
TABLE_SEP_RE = re.compile(r"^\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?$")
HR_RE = re.compile(r"^-{3,}\s*$")
SECTION_NO_RE = re.compile(r"^(\d+(?:\.\d+)*)\s+(.*)$")

INLINE_RE = re.compile(r"(\*\*.+?\*\*|`[^`]+`)")

BODY_FONT = "宋体"
HEAD_FONT = "微软雅黑"
HEAD_FONT_LATIN = "Arial"
MONO_FONT = "Consolas"

PAGE_CM_W, PAGE_CM_H = 21.0, 29.7
MARGIN_CM = 2.4
CONTENT_CM = PAGE_CM_W - 2 * MARGIN_CM  # 正文可用宽度 ≈ 16.2cm
MAX_IMG_CM_W, MAX_IMG_CM_H = CONTENT_CM, 20.0
IMG_DPI = 200  # PNG 像素 → 文档物理尺寸的换算密度

FIG_COLOR = RGBColor(0x59, 0x59, 0x59)


# --------------------------------------------------------------------------
# 1. mermaid → PNG（mermaid.ink）
# --------------------------------------------------------------------------

def render_mermaid(code: str, dest: Path) -> bytes:
    """经 mermaid.ink 将 mermaid 代码渲染为 PNG 并写入 dest。"""
    b64 = base64.urlsafe_b64encode(code.encode("utf-8")).decode("ascii")
    url = f"https://mermaid.ink/img/{b64}?type=png&bgColor=white"
    req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0 (build_tencent_docx)"})
    last_err: Exception | None = None
    for attempt in range(3):
        try:
            with urllib.request.urlopen(req, timeout=120) as resp:
                data = resp.read()
            if not data:
                raise ValueError("返回内容为空")
            dest.write_bytes(data)
            return data
        except Exception as exc:  # noqa: BLE001 - 网络重试
            last_err = exc
            print(f"    [重试 {attempt + 1}/3] {type(exc).__name__}: {str(exc)[:120]}")
    raise RuntimeError(f"渲染失败：{dest.name} -> {last_err}")


def validate_png(path: Path) -> tuple[int, int]:
    """校验 PNG 非空、magic bytes 正确、可打开且尺寸合理；返回 (宽, 高)。"""
    raw = path.read_bytes()
    if len(raw) < 1024:
        raise ValueError(f"{path.name} 体积过小（{len(raw)} bytes），疑似渲染失败")
    if raw[:8] != b"\x89PNG\r\n\x1a\n":
        raise ValueError(f"{path.name} 不是合法 PNG（magic bytes 不匹配）")
    with Image.open(io.BytesIO(raw)) as im:
        im.verify()  # 校验完整性
    with Image.open(io.BytesIO(raw)) as im:
        w, h = im.size
    if w < 100 or h < 60:
        raise ValueError(f"{path.name} 尺寸异常：{w}x{h}")
    return w, h


def prepare_diagrams(blocks: list[dict], force: bool = False) -> dict[str, tuple[int, int]]:
    """渲染/校验全部 mermaid 图，返回 {png 文件名: (宽, 高)}。"""
    DIAGRAM_DIR.mkdir(parents=True, exist_ok=True)
    sizes: dict[str, tuple[int, int]] = {}
    for idx, name in enumerate(DIAGRAM_NAMES):
        code = blocks[idx]["text"]
        dest = DIAGRAM_DIR / f"{name}.png"
        if force or not dest.exists():
            print(f"  渲染 {name}.png ...")
            render_mermaid(code, dest)
        w, h = validate_png(dest)
        sizes[f"{name}.png"] = (w, h)
        print(f"  校验 {name}.png OK  {w}x{h}px  {dest.stat().st_size} bytes")
    return sizes


# --------------------------------------------------------------------------
# 2. Markdown 解析
# --------------------------------------------------------------------------

def is_cjk(ch: str) -> bool:
    return ord(ch) > 0x2E7F


def smart_join(lines: list[str]) -> str:
    """中英混排软换行拼接：两侧均为 CJK 时不加空格，否则补一个空格。"""
    out = ""
    for line in lines:
        if out and not (is_cjk(out[-1]) and is_cjk(line[0])):
            out += " "
        out += line
    return out


def split_table_row(line: str) -> list[str]:
    s = line.strip()
    if s.startswith("|"):
        s = s[1:]
    if s.endswith("|"):
        s = s[:-1]
    return [c.strip() for c in s.split("|")]


def parse_blocks(lines: list[str]) -> list[dict]:
    """把 md 行序列解析为结构化 block 列表。"""
    blocks: list[dict] = []
    i, n = 0, len(lines)
    while i < n:
        line = lines[i]
        s = line.strip()
        if not s:
            i += 1
            continue

        # 围栏代码块 / mermaid
        if s.startswith("```"):
            lang = s[3:].strip()
            i += 1
            buf: list[str] = []
            while i < n and not lines[i].strip().startswith("```"):
                buf.append(lines[i])
                i += 1
            i += 1  # 跳过结束围栏
            kind = "mermaid" if lang == "mermaid" else "code"
            blocks.append({"kind": kind, "text": "\n".join(buf)})
            continue

        # 表格
        if s.startswith("|") and i + 1 < n and TABLE_SEP_RE.match(lines[i + 1].strip()):
            header = split_table_row(lines[i])
            i += 2
            rows = []
            while i < n and lines[i].strip().startswith("|"):
                rows.append(split_table_row(lines[i]))
                i += 1
            blocks.append({"kind": "table", "header": header, "rows": rows})
            continue

        # 分隔线
        if HR_RE.match(s):
            i += 1
            continue

        # 标题
        m = HEADING_RE.match(line)
        if m:
            blocks.append({"kind": "heading", "level": len(m.group(1)), "text": m.group(2).strip()})
            i += 1
            continue

        # 引用块：剥掉 '> ' 后递归解析（内部可能含列表/多段）
        if QUOTE_RE.match(line):
            buf = []
            while i < n and QUOTE_RE.match(lines[i]):
                buf.append(QUOTE_RE.match(lines[i]).group(1))
                i += 1
            blocks.extend(parse_blocks(buf))
            continue

        # 无序列表
        m = UL_RE.match(line)
        if m:
            blocks.append({"kind": "ul", "level": min(len(m.group(1)) // 2, 2), "text": m.group(2).strip()})
            i += 1
            continue

        # 有序列表
        m = OL_RE.match(line)
        if m:
            blocks.append({
                "kind": "ol",
                "level": min(len(m.group(1)) // 2, 2),
                "num": int(m.group(2)),
                "text": m.group(3).strip(),
            })
            i += 1
            continue

        # 段落（连续普通行合并）
        buf = []
        while i < n:
            cur = lines[i]
            cs = cur.strip()
            if not cs:
                break
            if (
                cs.startswith("```")
                or cs.startswith("|")
                or HR_RE.match(cs)
                or HEADING_RE.match(cur)
                or QUOTE_RE.match(cur)
                or UL_RE.match(cur)
                or OL_RE.match(cur)
            ):
                break
            buf.append(cs)
            i += 1
        blocks.append({"kind": "paragraph", "text": smart_join(buf)})
    return blocks


# --------------------------------------------------------------------------
# 3. docx 生成
# --------------------------------------------------------------------------

def _set_east_asia(style_or_run, east: str, latin: str | None = None) -> None:
    rpr = style_or_run.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:eastAsia"), east)
    if latin:
        rfonts.set(qn("w:ascii"), latin)
        rfonts.set(qn("w:hAnsi"), latin)


def setup_document(document: Document) -> None:
    sec = document.sections[0]
    sec.page_width = Cm(PAGE_CM_W)
    sec.page_height = Cm(PAGE_CM_H)
    for attr in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, attr, Cm(MARGIN_CM))

    normal = document.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(10.5)
    _set_east_asia(normal, BODY_FONT)
    normal.paragraph_format.line_spacing = 1.3
    normal.paragraph_format.space_after = Pt(4)

    sizes = {1: 20, 2: 16, 3: 13.5, 4: 12, 5: 11, 6: 10.5}
    for lvl, size in sizes.items():
        st = document.styles[f"Heading {lvl}"]
        st.font.name = HEAD_FONT_LATIN
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor(0x1F, 0x2A, 0x37)
        _set_east_asia(st, HEAD_FONT, HEAD_FONT_LATIN)
        st.paragraph_format.space_before = Pt(12 if lvl <= 2 else 8)
        st.paragraph_format.space_after = Pt(6)


class NumberingFactory:
    """为每个有序/无序列表创建独立编号定义（有序列表逐列表从 1 重新计数）。"""

    def __init__(self, document: Document) -> None:
        self.numbering = document.part.numbering_part.element
        self.abstract_id = self._max("w:abstractNum", "w:abstractNumId") + 1
        self.num_id = self._max("w:num", "w:numId") + 1

    def _max(self, tag: str, attr: str) -> int:
        ids = [int(e.get(qn(attr))) for e in self.numbering.findall(qn(tag))]
        return max(ids) if ids else 0

    def new_list(self, kind: str = "decimal") -> int:
        """kind: decimal（有序） | bullet（无序）。返回 numId。"""
        abstract_id, num_id = self.abstract_id, self.num_id
        self.abstract_id += 1
        self.num_id += 1

        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(abstract_id))
        multi = OxmlElement("w:multiLevelType")
        multi.set(qn("w:val"), "hybridMultilevel")
        abstract.append(multi)

        for lvl in range(3):
            lvl_el = OxmlElement("w:lvl")
            lvl_el.set(qn("w:ilvl"), str(lvl))
            start = OxmlElement("w:start")
            start.set(qn("w:val"), "1")
            lvl_el.append(start)
            fmt = OxmlElement("w:numFmt")
            fmt.set(qn("w:val"), kind)
            lvl_el.append(fmt)
            text = OxmlElement("w:lvlText")
            text.set(qn("w:val"), f"%{lvl + 1}." if kind == "decimal" else "•")
            lvl_el.append(text)
            jc = OxmlElement("w:lvlJc")
            jc.set(qn("w:val"), "left")
            lvl_el.append(jc)
            ppr = OxmlElement("w:pPr")
            ind = OxmlElement("w:ind")
            ind.set(qn("w:left"), str(480 + 360 * lvl))
            ind.set(qn("w:hanging"), "360")
            ppr.append(ind)
            lvl_el.append(ppr)
            abstract.append(lvl_el)

        first_num = self.numbering.find(qn("w:num"))
        if first_num is not None:
            first_num.addprevious(abstract)  # schema 要求 abstractNum 先于 num
        else:
            self.numbering.append(abstract)

        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(num_id))
        ref = OxmlElement("w:abstractNumId")
        ref.set(qn("w:val"), str(abstract_id))
        num.append(ref)
        self.numbering.append(num)
        return num_id


def apply_numbering(paragraph, num_id: int, ilvl: int) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl_el = OxmlElement("w:ilvl")
    ilvl_el.set(qn("w:val"), str(ilvl))
    num_pr.append(ilvl_el)
    num_id_el = OxmlElement("w:numId")
    num_id_el.set(qn("w:val"), str(num_id))
    num_pr.append(num_id_el)
    pstyle = ppr.find(qn("w:pStyle"))
    if pstyle is not None:
        pstyle.addnext(num_pr)
    else:
        ppr.insert(0, num_pr)


def add_inline_runs(paragraph, text: str, *, bold: bool = False, size: float | None = None,
                    mono: bool = False, color: RGBColor | None = None) -> None:
    """写入行内内容，处理 **加粗** 与 `行内代码`。"""
    for token in INLINE_RE.split(text):
        if not token:
            continue
        is_bold, is_code = bold, mono
        content = token
        if token.startswith("**") and token.endswith("**") and len(token) > 4:
            is_bold, content = True, token[2:-2]
        elif token.startswith("`") and token.endswith("`") and len(token) >= 2:
            is_code, content = True, token[1:-1]
        run = paragraph.add_run(content)
        run.bold = is_bold
        if is_code:
            run.font.name = MONO_FONT
            _set_east_asia(run, BODY_FONT, MONO_FONT)
            run.font.size = Pt(size - 0.5) if size else Pt(9.5)
        elif size:
            run.font.size = Pt(size)
        if color is not None:
            run.font.color.rgb = color


def content_visual_len(text: str) -> int:
    """粗略视觉宽度：CJK 记 2，其余记 1（用于列宽分配）。"""
    return sum(2 if is_cjk(c) else 1 for c in text)


def add_table(document: Document, header: list[str], rows: list[list[str]]) -> None:
    n_cols = len(header)
    table = document.add_table(rows=1, cols=n_cols)
    table.style = "Table Grid"
    table.autofit = False

    hdr_cells = table.rows[0].cells
    for cell, text in zip(hdr_cells, header):
        cell.paragraphs[0].paragraph_format.space_after = Pt(2)
        add_inline_runs(cell.paragraphs[0], text, bold=True, size=9.5)
    # 表头跨页重复
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)

    for row in rows:
        cells = table.add_row().cells
        for cell, text in zip(cells, (row + [""] * n_cols)[:n_cols]):
            cell.paragraphs[0].paragraph_format.space_after = Pt(2)
            add_inline_runs(cell.paragraphs[0], text, size=9.5)

    # 按内容宽度比例分配列宽（带最小宽度约束）
    weights = []
    for idx in range(n_cols):
        col_texts = [header[idx]] + [r[idx] if idx < len(r) else "" for r in rows]
        weights.append(max(content_visual_len(t) for t in col_texts))
    min_w = 1.7
    total = CONTENT_CM
    avail = total - min_w * n_cols
    wsum = sum(weights) or 1
    widths = [min_w + avail * w / wsum for w in weights]
    for row in table.rows:
        for cell, w in zip(row.cells, widths):
            cell.width = Cm(round(w, 2))


def add_code_block(document: Document, text: str) -> None:
    """代码块 → 等宽纯文本段落（保留缩进与换行）。"""
    for line in text.split("\n"):
        p = document.add_paragraph()
        pf = p.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.line_spacing = 1.0
        run = p.add_run(line if line.strip() else " ")
        run.font.name = MONO_FONT
        _set_east_asia(run, BODY_FONT, MONO_FONT)
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x24, 0x29, 0x2E)


def add_figure(document: Document, png: Path, px_size: tuple[int, int], caption: str) -> None:
    px_w, px_h = px_size
    w_cm, h_cm = px_w / IMG_DPI * 2.54, px_h / IMG_DPI * 2.54
    k = min(MAX_IMG_CM_W / w_cm, MAX_IMG_CM_H / h_cm)
    if k < 1:
        w_cm, h_cm = w_cm * k, h_cm * k
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(str(png), width=Cm(round(w_cm, 2)), height=Cm(round(h_cm, 2)))

    cap = document.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(10)
    run = cap.add_run(caption)
    run.font.size = Pt(9)
    run.font.color.rgb = FIG_COLOR
    _set_east_asia(run, BODY_FONT)


def build_docx(blocks: list[dict], sizes: dict[str, tuple[int, int]], out_path: Path) -> None:
    document = Document()
    setup_document(document)
    try:
        numbering: NumberingFactory | None = NumberingFactory(document)
    except Exception as exc:  # noqa: BLE001 - 极端情况下降级为手工编号
        print(f"  [警告] 编号定义不可用（{exc}），降级为手写序号")
        numbering = None

    current_heading = ""
    section_no = ""
    heading_name = ""
    preceding_label = ""
    fig_no = 0
    cur_num_id: int | None = None
    prev_kind: str | None = None

    for blk in blocks:
        kind = blk["kind"]

        if kind == "heading":
            text = blk["text"]
            m = SECTION_NO_RE.match(text)
            section_no = m.group(1) if m else ""
            name = (m.group(2) if m else text).replace("【", "（").replace("】", "）")
            heading_name = name
            current_heading = text
            lvl = min(blk["level"], 6)
            document.add_heading(text, level=lvl)
            preceding_label = ""
            cur_num_id = None

        elif kind == "paragraph":
            p = document.add_paragraph()
            add_inline_runs(p, blk["text"])
            stripped = blk["text"].strip()
            if re.fullmatch(r"\*\*.+?\*\*[：:]?", stripped):
                preceding_label = stripped.strip("*：: ")
            else:
                preceding_label = ""
            cur_num_id = None

        elif kind == "table":
            add_table(document, blk["header"], blk["rows"])
            document.add_paragraph().paragraph_format.space_after = Pt(0)
            cur_num_id = None

        elif kind == "code":
            add_code_block(document, blk["text"])
            add_blank_spacer(document)
            preceding_label = ""
            cur_num_id = None

        elif kind == "mermaid":
            fig_no += 1
            png = DIAGRAM_DIR / f"{DIAGRAM_NAMES[fig_no - 1]}.png"
            parts = [f"图 {fig_no}"]
            if section_no:
                parts.append(f"§{section_no}")
            if heading_name:
                parts.append(heading_name)
            caption = "　".join(parts[:2]) + ("　" + "｜".join(parts[2:]) if len(parts) > 2 else "")
            if preceding_label:
                caption += f"｜{preceding_label}"
            add_figure(document, png, sizes[png.name], caption)
            preceding_label = ""
            cur_num_id = None

        elif kind in ("ul", "ol"):
            level = blk["level"]
            if kind == "ol":
                if prev_kind != "ol" or cur_num_id is None:
                    cur_num_id = numbering.new_list("decimal") if numbering else None
                num_id = cur_num_id
            else:
                num_id = numbering.new_list("bullet") if numbering else None
            p = document.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            text = blk["text"]
            if num_id is not None:
                apply_numbering(p, num_id, level)
            elif kind == "ol":
                text = f"{blk['num']}. {text}"
            add_inline_runs(p, text)
            if kind == "ul":
                cur_num_id = None
            preceding_label = ""

        prev_kind = kind

    document.save(str(out_path))
    print(f"已生成：{out_path}")


def add_blank_spacer(document: Document) -> None:
    p = document.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    run = p.add_run(" ")
    run.font.size = Pt(4)


# --------------------------------------------------------------------------
# 4. 校验
# --------------------------------------------------------------------------

def verify_docx(path: Path, expected_images: int | None = None) -> None:
    expected = len(DIAGRAM_NAMES) if expected_images is None else expected_images
    document = Document(str(path))
    n_images = len(document.inline_shapes)
    n_tables = len(document.tables)
    heads = [(p.style.name, p.text) for p in document.paragraphs if p.style.name.startswith("Heading")]
    h_counts: dict[str, int] = {}
    for style, _ in heads:
        h_counts[style] = h_counts.get(style, 0) + 1
    print("\n=== 校验结果 ===")
    print(f"文件：{path}")
    print(f"图片数：{n_images}（期望 {expected}）")
    print(f"表格数：{n_tables}")
    print(f"标题数：{len(heads)}  {dict(sorted(h_counts.items()))}")
    print("一级/二级标题：")
    for style, text in heads:
        if style in ("Heading 1", "Heading 2"):
            print(f"  [{style}] {text}")
    assert n_images == expected, f"图片数不符：{n_images} != {expected}"
    print("校验通过。")


# --------------------------------------------------------------------------

def main() -> None:
    parser = argparse.ArgumentParser(description="6AM md → 腾讯文档可导入 docx")
    parser.add_argument("--force-render", action="store_true", help="强制重新渲染全部 mermaid PNG")
    parser.add_argument("--src", type=Path, default=SRC_MD, help="输入 markdown 路径（默认原版）")
    parser.add_argument("--out", type=Path, default=OUT_DOCX, help="输出 docx 路径（默认原版）")
    args = parser.parse_args()

    md_text = args.src.read_text(encoding="utf-8")
    mermaid_blocks = [{"kind": "mermaid", "text": m} for m in MERMAID_RE.findall(md_text)]
    print(f"从 {args.src.name} 提取 mermaid 代码块：{len(mermaid_blocks)} 个")
    if mermaid_blocks and len(mermaid_blocks) != len(DIAGRAM_NAMES):
        sys.exit(f"mermaid 块数量异常：{len(mermaid_blocks)} != {len(DIAGRAM_NAMES)}")

    print("\n[1/3] 渲染 / 校验 mermaid PNG")
    sizes = prepare_diagrams(mermaid_blocks, force=args.force_render) if mermaid_blocks else {}

    print("\n[2/3] 解析 markdown 并生成 docx")
    blocks = parse_blocks(md_text.split("\n"))
    kinds: dict[str, int] = {}
    for b in blocks:
        kinds[b["kind"]] = kinds.get(b["kind"], 0) + 1
    print(f"  block 统计：{kinds}")
    build_docx(blocks, sizes, args.out)

    print("\n[3/3] 校验 docx")
    verify_docx(args.out, expected_images=len(mermaid_blocks))


if __name__ == "__main__":
    main()