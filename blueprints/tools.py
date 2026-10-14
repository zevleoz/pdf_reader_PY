"""工具蓝图：预约管理、解读会纪要、内部下载、Prompt Lab。"""
from __future__ import annotations

import json
import os
from datetime import date, timedelta
from typing import Any, Dict, List

from flask import Blueprint, jsonify, render_template, request, send_file

import db as _db
import extract
import evaluation_rules as _eval_rules
from minutes_schema import validate_minutes, minutes_to_markdown
from minutes_docx import build_minutes_docx
from app_utils import (BASE_DIR, DATA_DIR, admin_required, role_required,
                       student_scoped, minutes_scoped)
from ai_client import (_assemble_e4_protocol, _build_phase1_md_text,
                       _build_qualification_view, _build_y4_interpret_context,
                       _build_y4_payload, _dashscope_chat,
                       _determine_list_placement, _e4_step2_system,
                       _e4_step2_tail, _e4_step3_system,
                       _finalize_y4_interpretation, _format_e4_line,
                       _get_student_name_from_items, _sanitize_e4_step2)

bp = Blueprint("tools", __name__)


@bp.route("/transcript")
@role_required('admin', 'advisor')
def transcript_page():
    """Render the transcript upload / summary generation page."""
    return render_template("transcript.html")


@bp.route("/api/transcript", methods=["POST"])
@admin_required
def api_transcript():
    import urllib.request as _ureq

    data = request.get_json(force=True)
    transcript_text = (data.get("transcript") or "").strip()
    student_name = (data.get("student_name") or "").strip()
    student_grade = (data.get("student_grade") or "").strip()
    student_gender = (data.get("student_gender") or "").strip()
    student_id = data.get("student_id")
    report_id = data.get("report_id")
    custom_prompt = (data.get("custom_prompt") or "").strip()

    if len(transcript_text) < 50:
        return jsonify({"ok": False, "error": "逐字稿太短，至少 50 字"}), 400

    # ------------------------------------------------------------------
    # 1) Resolve student + report info (if student_id/report_id given)
    # ------------------------------------------------------------------
    resolved_sid = None
    resolved_rid = None
    if student_id:
        try:
            resolved_sid = int(student_id)
        except (TypeError, ValueError):
            return jsonify({"ok": False, "error": "student_id 格式错误"}), 400
    if report_id:
        try:
            resolved_rid = int(report_id)
        except (TypeError, ValueError):
            return jsonify({"ok": False, "error": "report_id 格式错误"}), 400

    report_summary = "（暂无测评数据）"

    if resolved_rid:
        # Specific report given → use its raw data
        rec = _db.get_report_raw(resolved_rid)
        if rec:
            raw = rec.get("raw") or {}
            schema_items = raw.get("schema_124", [])
            lines = []
            for item in schema_items:
                val = item.get("value", "")
                if val and str(val).strip() not in ("", "—", None):
                    lines.append(f"  {item.get('label', item.get('code', '?'))}: {val}")
            if lines:
                report_summary = "\n".join(lines[:60])
            # Fill empty student info from report record
            if not student_name:
                student_name = rec.get("student_name", "")
            stu_info = (raw.get("student") or {}) if isinstance(raw, dict) else {}
            if not student_grade:
                student_grade = rec.get("grade") or stu_info.get("grade", "") or ""
            if not student_gender:
                student_gender = stu_info.get("gender", "") or ""
    elif resolved_sid:
        # Student only, no specific report → take the latest report for that student
        reports = _db.get_student_reports(resolved_sid)
        if reports:
            latest = reports[0]
            resolved_rid = latest.get("id")
            rec = _db.get_report_raw(resolved_rid)
            if rec:
                raw = rec.get("raw") or {}
                schema_items = raw.get("schema_124", [])
                lines = []
                for item in schema_items:
                    val = item.get("value", "")
                    if val and str(val).strip() not in ("", "—", None):
                        lines.append(f"  {item.get('label', item.get('code', '?'))}: {val}")
                if lines:
                    report_summary = "\n".join(lines[:60])
                if not student_name:
                    student_name = rec.get("student_name", "")
                stu_info = (raw.get("student") or {}) if isinstance(raw, dict) else {}
                if not student_grade:
                    student_grade = rec.get("grade") or stu_info.get("grade", "") or ""
                if not student_gender:
                    student_gender = stu_info.get("gender", "") or ""
        # If student has no reports, fallback to name from students table
        if not student_name:
            for s in _db.get_students():
                if s["id"] == resolved_sid:
                    student_name = s.get("name", "")
                    student_grade = student_grade or s.get("grade", "") or ""
                    break
    else:
        # Backward-compat: read current-session report_data.json
        report_path = DATA_DIR / "report_data.json"
        if report_path.exists():
            try:
                rd = json.loads(report_path.read_text(encoding="utf-8"))
                s124 = rd.get("schema_124", [])
                lines = []
                for item in s124:
                    val = item.get("value", "")
                    if val and str(val).strip() not in ("", "—", None):
                        lines.append(f"  {item.get('label', item.get('code', '?'))}: {val}")
                if lines:
                    report_summary = "\n".join(lines[:60])
                # Also autofill student info
                stu_info = rd.get("student", {}) or {}
                if not student_name:
                    student_name = stu_info.get("name", "")
                if not student_grade:
                    student_grade = stu_info.get("grade", "") or ""
                if not student_gender:
                    student_gender = stu_info.get("gender", "") or ""
            except Exception:
                pass

    # ------------------------------------------------------------------
    # 2) Build system prompt (custom_prompt overrides file)
    # ------------------------------------------------------------------
    # structured_mode: 无 custom_prompt 且 v2 结构化 Prompt 存在时，走 JSON 输出管线。
    structured_mode = False
    if custom_prompt:
        system_prompt = custom_prompt
    else:
        v2_path = BASE_DIR / "prompts" / "transcript_summary_v2.md"
        if v2_path.exists():
            system_prompt = v2_path.read_text(encoding="utf-8")
            structured_mode = True
        else:
            prompt_path = BASE_DIR / "prompts" / "transcript_summary.md"
            system_prompt = prompt_path.read_text(encoding="utf-8") if prompt_path.exists() else "你是解读会纪要撰写人。"

    # No string-replace on system_prompt. All dynamic context goes in user
    # message to match the user prompt's framing ("基于接下来提供的...").

    # ------------------------------------------------------------------
    # 3) Assemble messages
    # ------------------------------------------------------------------
    effective_name = student_name or "未填写"
    user_content = (
        "## 学生基本信息\n"
        f"姓名：{effective_name}\n"
        f"年级：{student_grade or '—'}\n"
        f"性别：{student_gender or '—'}\n\n"
        "## Y4 测评数据摘要（供交叉验证参考，如果和会议内容无关可以忽略）\n"
        f"{report_summary}\n\n"
        "## 解读会录音转写/会议纪要\n"
        f"{transcript_text[:12000]}\n\n"
        "请基于以上提供的录音转写/会议纪要，直接生成最终会议记录。"
    )

    messages = [
        {"role": "system", "content": system_prompt},
        {"role": "user", "content": user_content},
    ]

    # ------------------------------------------------------------------
    # 4) Call DashScope API — prefer ITERATION_MODEL (qwen-turbo by default)
    # ------------------------------------------------------------------
    dashscope_key = os.environ.get("DASHSCOPE_API_KEY", extract.DEFAULT_DASHSCOPE_KEY).strip()
    model_name = (
        os.environ.get("ITERATION_MODEL")
        or os.environ.get("AI_TEXT_MODEL")
        or "qwen-turbo"
    )
    url = "https://dashscope.aliyuncs.com/compatible-mode/v1/chat/completions"
    req_body = {
        "model": model_name,
        "messages": messages,
        "temperature": 0.7,
        "max_tokens": 8192,
    }
    if structured_mode:
        req_body["response_format"] = {"type": "json_object"}
    payload = json.dumps(req_body).encode("utf-8")
    req = _ureq.Request(
        url, data=payload,
        headers={"Authorization": f"Bearer {dashscope_key}",
                 "Content-Type": "application/json"},
        method="POST",
    )
    try:
        with _ureq.urlopen(req, timeout=180) as resp:
            result = json.loads(resp.read().decode("utf-8"))
        raw_content = result["choices"][0]["message"]["content"]
    except Exception as exc:
        return jsonify({"ok": False, "error": f"AI 调用失败: {exc}"}), 500

    # ------------------------------------------------------------------
    # 5) Structured JSON → normalize → Markdown preview; 失败回退 legacy
    # ------------------------------------------------------------------
    minutes_json_str = None
    if structured_mode:
        try:
            parsed = json.loads(raw_content)
            d = validate_minutes(parsed)
            # 请求/报告解析出的学生信息优先于 LLM 抽取结果
            smeta = d["meeting_meta"]
            if student_name:
                smeta["student_name"] = student_name
            if student_grade:
                smeta["grade"] = student_grade
            if student_gender:
                smeta["gender"] = student_gender
            summary = minutes_to_markdown(d)
            minutes_json_str = json.dumps(d, ensure_ascii=False)
        except Exception as exc:
            print(f"[transcript] 结构化解析失败，回退 legacy: {exc}")
            summary = raw_content
    else:
        summary = raw_content

    # ------------------------------------------------------------------
    # 6) Fix title placeholder: 【学生姓名】 → actual name (legacy 路径)
    # ------------------------------------------------------------------
    if effective_name and effective_name != "未填写":
        summary = summary.replace("【学生姓名】", effective_name)

    # ------------------------------------------------------------------
    # 7) Persist to DB if student_id is linked
    # ------------------------------------------------------------------
    minutes_id = None
    if resolved_sid:
        try:
            minutes_id = _db.add_minutes(
                student_id=resolved_sid,
                report_id=resolved_rid,
                transcript_text=transcript_text[:50000],
                minutes_text=summary,
                minutes_json=minutes_json_str,
            )
        except Exception as e:
            print(f"[DB] 会议纪要保存失败 (非致命): {e}")

    return jsonify({
        "ok": True,
        "summary": summary,
        "minutes_id": minutes_id,
        "student_id": resolved_sid,
        "report_id": resolved_rid,
        "student_name": student_name,
        "student_grade": student_grade,
        "student_gender": student_gender,
    })


def _build_docx(summary_text: str, student_name: str = "", student_grade: str = "",
               student_gender: str = "") -> bytes:
    """Convert markdown-style meeting minutes text into a styled Word .docx file.
    Supports #/##/### ATX headings, **bold**, bullet lines (- / •), [ ]/[x] checkboxes,
    and quote lines. Preserves the old **一、xxx** bold-heading format for backward compat.
    Returns the raw bytes of the docx file.
    """
    import re
    import io
    from docx import Document
    from docx.shared import Pt, RGBColor, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()

    # ---- Page margins ----
    for section in doc.sections:
        section.top_margin = Cm(2.2)
        section.bottom_margin = Cm(2.2)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    # ---- Base styles ----
    style = doc.styles["Normal"]
    style.font.name = "微软雅黑"
    style.font.size = Pt(11)
    style.font.color.rgb = RGBColor(0x14, 0x14, 0x14)

    # ---- Title ----
    title_line = "凭远教育 · Y4 解读会会议纪要"
    if student_name:
        title_line += f" — {student_name}"
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run(title_line)
    title_run.bold = True
    title_run.font.size = Pt(18)
    title_run.font.color.rgb = RGBColor(0x14, 0x14, 0x14)
    title_run.font.name = "微软雅黑"
    title_p.paragraph_format.space_after = Pt(6)

    # ---- Student info table ----
    from datetime import datetime
    info_line = f"学生：{student_name or '—'}"
    if student_grade:
        info_line += f"　｜　年级：{student_grade}"
    if student_gender:
        info_line += f"　｜　性别：{student_gender}"
    info_line += f"　｜　生成时间：{datetime.now().strftime('%Y-%m-%d')}"
    meta_p = doc.add_paragraph()
    meta_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta_run = meta_p.add_run(info_line)
    meta_run.font.size = Pt(10)
    meta_run.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)
    meta_p.paragraph_format.space_after = Pt(16)

    # ---- Red divider ----
    div_p = doc.add_paragraph()
    div_run = div_p.add_run("━" * 30)
    div_run.font.color.rgb = RGBColor(0xB3, 0x3A, 0x3A)
    div_run.font.size = Pt(10)
    div_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    div_p.paragraph_format.space_after = Pt(12)

    # ---- Parse lines and build document structure ----
    # Heading patterns: **一、** legacy bold-heading; ATX # / ## / ###
    heading_legacy_re = re.compile(r"^\s*\*\*(.+?)\*\*\s*(?:（(.+?)）)?\s*$")
    atx_h1_re = re.compile(r"^\s*#\s+(.+?)\s*$")
    atx_h2_re = re.compile(r"^\s*##\s+(.+?)\s*$")
    atx_h3_re = re.compile(r"^\s*###\s+(.+?)\s*$")
    # Bullet markers
    bullet_chars = ("•", "·", "-", "*", "【")

    lines = summary_text.split("\n")

    for raw in lines:
        line = raw.rstrip()
        stripped = line.strip()

        # Skip separator lines like "---"
        if re.match(r"^-+$", stripped) or re.match(r"^=+$", stripped) or re.match(r"^━+$", stripped):
            continue
        if stripped.startswith("> ") or stripped == ">":
            # Markdown quote line → treat as italic small paragraph
            if stripped == ">":
                continue
            q = stripped[2:].strip()
            if q:
                p = doc.add_paragraph()
                run = p.add_run(q)
                run.italic = True
                run.font.size = Pt(10)
                run.font.color.rgb = RGBColor(0x4A, 0x4A, 0x4A)
            continue

        # === ATX headings (new prompt format, check first so # wins over other patterns) ===
        m_h1 = atx_h1_re.match(stripped)
        if m_h1:
            hp = doc.add_paragraph()
            r = hp.add_run(m_h1.group(1).strip())
            r.bold = True
            r.font.size = Pt(16)
            r.font.color.rgb = RGBColor(0xB3, 0x3A, 0x3A)  # red (H1)
            r.font.name = "微软雅黑"
            hp.paragraph_format.space_before = Pt(14)
            hp.paragraph_format.space_after = Pt(6)
            continue
        m_h2 = atx_h2_re.match(stripped)
        if m_h2:
            hp = doc.add_paragraph()
            r = hp.add_run(m_h2.group(1).strip())
            r.bold = True
            r.font.size = Pt(14)
            r.font.color.rgb = RGBColor(0xB3, 0x3A, 0x3A)  # deep red (H2)
            r.font.name = "微软雅黑"
            hp.paragraph_format.space_before = Pt(12)
            hp.paragraph_format.space_after = Pt(4)
            continue
        m_h3 = atx_h3_re.match(stripped)
        if m_h3:
            hp = doc.add_paragraph()
            r = hp.add_run(m_h3.group(1).strip())
            r.bold = True
            r.font.size = Pt(12)
            r.font.color.rgb = RGBColor(0x14, 0x14, 0x14)  # ink (H3)
            r.font.name = "微软雅黑"
            hp.paragraph_format.space_before = Pt(10)
            hp.paragraph_format.space_after = Pt(3)
            continue

        # Legacy heading: **一、核心发现 · 3 条**（...）
        m_leg = heading_legacy_re.match(stripped)
        if m_leg:
            section_title = m_leg.group(1).strip()
            section_subtitle = m_leg.group(2).strip() if m_leg.group(2) else None
            hp = doc.add_paragraph()
            hrun = hp.add_run(section_title)
            hrun.bold = True
            hrun.font.size = Pt(14)
            hrun.font.color.rgb = RGBColor(0xB3, 0x3A, 0x3A)
            hrun.font.name = "微软雅黑"
            hp.paragraph_format.space_before = Pt(12)
            hp.paragraph_format.space_after = Pt(4)
            if section_subtitle:
                subrun = hp.add_run(f" （{section_subtitle}）")
                subrun.bold = False
                subrun.font.size = Pt(10)
                subrun.font.color.rgb = RGBColor(0x8A, 0x8A, 0x8A)
            continue

        if not stripped:
            continue  # empty line

        # Check if it's a sub-heading line (bold **...** without a bullet)
        if stripped.startswith("**") and stripped.endswith("**") and not any(stripped.startswith(b) for b in bullet_chars):
            subh = stripped[2:-2]
            hp2 = doc.add_paragraph()
            r2 = hp2.add_run(subh)
            r2.bold = True
            r2.font.size = Pt(11.5)
            r2.font.color.rgb = RGBColor(0x14, 0x14, 0x14)
            hp2.paragraph_format.space_before = Pt(8)
            hp2.paragraph_format.space_after = Pt(2)
            continue

        # Bullet lines (starts with bullet char, possibly after whitespace)
        is_bullet = False
        content = stripped
        checkbox = ""
        # Strip leading bullet marker
        lead = re.match(r"^(\s*[-•·*·]\s*(?:\[\s*[xX\s]\]\s*)?)", content)
        if lead:
            marker = lead.group(1)
            if "[ ]" in marker or "［］" in marker or "[  ]" in marker:
                checkbox = "☐"
            elif re.search(r"\[\s*[xX]\s*\]", marker):
                checkbox = "☑"
            content = content[lead.end():].strip()
            is_bullet = True

        if not is_bullet and (content.startswith("【") or content.startswith("• ") or content.startswith("· ")):
            is_bullet = True
            if content.startswith("• "):
                content = content[2:]
            elif content.startswith("· "):
                content = content[2:]

        # Strip inline **...** bold markers and convert to rich runs
        tokens = re.split(r"(\*\*[^*]+\*\*)", content)
        if is_bullet:
            bullet_prefix = f"{checkbox} • " if checkbox else "• "
            # Build paragraph with bullet marker (no List Bullet style for consistency)
            p2 = doc.add_paragraph()
            b_run = p2.add_run(bullet_prefix)
            b_run.font.color.rgb = RGBColor(0xB3, 0x3A, 0x3A)
            b_run.bold = True
            p2.paragraph_format.space_before = Pt(2)
            p2.paragraph_format.space_after = Pt(2)
            p2.paragraph_format.left_indent = Cm(0.6)
            # Append content tokens
            for tok in tokens:
                if tok.startswith("**") and tok.endswith("**"):
                    r = p2.add_run(tok[2:-2])
                    r.bold = True
                    r.font.size = Pt(11)
                else:
                    r = p2.add_run(tok)
                    r.font.size = Pt(11)
            continue

        # Plain paragraph (non-heading, non-bullet)
        pp = doc.add_paragraph()
        pp.paragraph_format.space_before = Pt(2)
        pp.paragraph_format.space_after = Pt(2)
        pp.paragraph_format.line_spacing = Pt(20)
        for tok in tokens:
            if tok.startswith("**") and tok.endswith("**"):
                r = pp.add_run(tok[2:-2])
                r.bold = True
                r.font.size = Pt(11)
            else:
                r = pp.add_run(tok)
                r.font.size = Pt(11)

    # ---- Footer divider ----
    doc.add_paragraph()
    footer_div = doc.add_paragraph()
    fd_run = footer_div.add_run("— 凭远教育 · Y4 综合测评系统 —")
    fd_run.font.size = Pt(9)
    fd_run.font.color.rgb = RGBColor(0x8A, 0x8A, 0x8A)
    footer_div.alignment = WD_ALIGN_PARAGRAPH.CENTER

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


@bp.route("/api/transcript/docx", methods=["POST"])
@admin_required
def api_transcript_docx():
    """Generate minutes as DOCX file.

    Two input modes:
      A) {minutes_id: N} → load from DB, auto-fills student info + minutes_text
      B) {summary, student_name, student_grade, student_gender} → legacy mode

    Returns the .docx binary as attachment download.
    """
    data = request.get_json(force=True)
    minutes_id = data.get("minutes_id")
    summary_text = ""
    student_name = ""
    student_grade = ""
    student_gender = ""
    minutes_json = None

    if minutes_id:
        mrec = _db.get_minutes(int(minutes_id))
        if not mrec:
            return jsonify({"ok": False, "error": "纪要不存在"}), 404
        summary_text = mrec.get("minutes_text", "") or ""
        student_name = mrec.get("student_name", "") or ""
        student_grade = mrec.get("student_grade", "") or ""
        student_gender = mrec.get("student_gender", "") or ""
        minutes_json = mrec.get("minutes_json")
    else:
        summary_text = (data.get("summary") or "").strip()
        student_name = (data.get("student_name") or "").strip()
        student_grade = (data.get("student_grade") or "").strip()
        student_gender = (data.get("student_gender") or "").strip()

    # 结构化纪要优先走模板渲染；无 minutes_json 则回退旧渲染器（含手动编辑后的记录）
    if minutes_json:
        try:
            d = validate_minutes(json.loads(minutes_json))
            smeta = d["meeting_meta"]
            if student_name:
                smeta["student_name"] = student_name
            if student_grade:
                smeta["grade"] = student_grade
            if student_gender:
                smeta["gender"] = student_gender
            docx_bytes = build_minutes_docx(d)
        except Exception as exc:
            print(f"[transcript/docx] 结构化渲染失败，回退 legacy: {exc}")
            docx_bytes = None
    else:
        docx_bytes = None

    if docx_bytes is None:
        if not summary_text:
            return jsonify({"ok": False, "error": "纪要内容为空"}), 400
        try:
            docx_bytes = _build_docx(summary_text, student_name, student_grade, student_gender)
        except Exception as exc:
            import traceback
            tb = traceback.format_exc()
            return jsonify({"ok": False, "error": f"DOCX 生成失败: {exc}", "trace": tb}), 500

    filename_suffix = student_name if student_name else "纪要"
    filename = f"凭远Y4解读会纪要_{filename_suffix}.docx"
    safe_name = filename.encode("utf-8").decode("latin-1", "ignore")
    from urllib.parse import quote
    from flask import Response
    resp = Response(
        docx_bytes,
        mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={
            # filename* 必须为纯 ASCII（RFC 5987 percent-encoding），否则 WSGI latin-1 编码报错
            "Content-Disposition": f"attachment; filename=\"{safe_name}\"; filename*=UTF-8''{quote(filename)}"
        }
    )
    return resp


# ---------------------------------------------------------------------------
# Meeting minutes management APIs
# ---------------------------------------------------------------------------

@bp.route("/api/students/<int:student_id>/minutes")
@admin_required
@student_scoped
def api_student_minutes(student_id):
    """List all historical meeting minutes for a student (desc by created_at)."""
    items = _db.get_minutes_by_student(student_id)
    return jsonify({"ok": True, "minutes": items})


@bp.route("/api/minutes/<int:minutes_id>")
@admin_required
@minutes_scoped
def api_minutes_detail(minutes_id):
    """Get single meeting minutes: transcript + minutes_text + student info."""
    item = _db.get_minutes(minutes_id)
    if not item:
        return jsonify({"ok": False, "error": "纪要不存在"}), 404
    return jsonify({"ok": True, **item})


@bp.route("/api/minutes/<int:minutes_id>", methods=["PUT"])
@admin_required
@minutes_scoped
def api_minutes_update(minutes_id):
    """Update minutes (e.g. re-generate with new transcript / manual edits)."""
    data = request.get_json(force=True) or {}
    fields = {}
    for k in ("transcript_text", "minutes_text", "report_id"):
        if k in data and data[k] is not None:
            fields[k] = data[k]
    # 手动编辑 minutes_text 会使其结构化版本过期 → 置空 minutes_json，DOCX 回退 legacy 渲染
    if "minutes_text" in fields:
        fields["minutes_json"] = None
    if not fields:
        return jsonify({"ok": True})
    try:
        _db.update_minutes(minutes_id, **fields)
    except Exception as e:
        return jsonify({"ok": False, "error": str(e)}), 500
    return jsonify({"ok": True})


@bp.route("/api/minutes/<int:minutes_id>", methods=["DELETE"])
@admin_required
@minutes_scoped
def api_minutes_delete(minutes_id):
    try:
        _db.delete_minutes(minutes_id)
    except Exception as e:
        return jsonify({"ok": False, "error": str(e)}), 500
    return jsonify({"ok": True})
@bp.route("/booking")
def booking_page():
    return render_template("booking.html")


@bp.route("/api/booking", methods=["POST"])
def api_booking():
    from datetime import datetime
    data = request.get_json(force=True)
    name = (data.get("student_name") or "").strip()
    if not name:
        return jsonify({"ok": False, "error": "学生姓名必填"}), 400
    appt_time_str = data.get("appointment_time", "")
    try:
        appt_time = datetime.fromisoformat(appt_time_str)
    except ValueError:
        return jsonify({"ok": False, "error": "时间格式错误"}), 400

    # Check slot capacity
    booking_date = appt_time.date()
    time_slot_str = appt_time.strftime("%H:%M")
    booking_counts = _db.get_slot_booking_counts(booking_date)
    booked = booking_counts.get(time_slot_str, 0)
    if booked >= SLOT_CAPACITY:
        return jsonify({"ok": False, "error": "该时段已满（" + str(booked) + "/" + str(SLOT_CAPACITY) + "），请选择其他时间"}), 409

    # Check admin hasn't closed this slot
    avail_entries = _db.get_availability(booking_date)
    avail_map = {e["time_slot"]: e["is_available"] for e in avail_entries}
    is_open = avail_map.get(time_slot_str, True)
    if not is_open:
        return jsonify({"ok": False, "error": "该时段已关闭，请选择其他时间"}), 409

    # Create student + booking in one transaction (auto-archive)
    student_id, booking_id = _db.create_booking_with_student(
        student_name=name,
        appointment_time=appt_time,
        advisor_name=(data.get("advisor_name") or "").strip(),
        school=(data.get("school") or "").strip(),
        single_parent=data.get("single_parent", "false"),
        notes=data.get("notes", ""),
    )
    return jsonify({"ok": True, "booking_id": booking_id, "student_id": student_id})


@bp.route("/admin/bookings")
@role_required('admin', 'advisor')
def admin_bookings_page():
    return render_template("admin_bookings.html")


@bp.route("/api/bookings")
@admin_required
def api_bookings():
    status = request.args.get("status")
    bookings = _db.get_bookings(status=status)
    return jsonify({"ok": True, "bookings": bookings})


@bp.route("/api/booking/<int:booking_id>/complete", methods=["POST"])
@admin_required
def api_booking_complete(booking_id):
    try:
        sid = _db.complete_booking(booking_id)
        return jsonify({"ok": True, "student_id": sid})
    except ValueError as e:
        return jsonify({"ok": False, "error": str(e)}), 404


@bp.route("/api/booking/<int:booking_id>/cancel", methods=["POST"])
@admin_required
def api_booking_cancel(booking_id):
    _db.update_booking_status(booking_id, "cancelled")
    return jsonify({"ok": True})


# ---------------------------------------------------------------------------
# Availability management
# ---------------------------------------------------------------------------
SLOT_CAPACITY = 4  # Max students per time slot


@bp.route("/api/availability")
def api_get_availability():
    """Get availability for a given date. Query param: date=YYYY-MM-DD
    Returns each slot with is_available, booked_count, and capacity.
    Default: all slots available. Admin can close slots. Full slots (4/4) are unavailable.
    """
    date_str = request.args.get("date", "")
    if not date_str:
        return jsonify({"ok": False, "error": "请指定日期"}), 400
    try:
        date_val = date.fromisoformat(date_str)
    except ValueError:
        return jsonify({"ok": False, "error": "日期格式错误"}), 400
    entries = _db.get_availability(date_val)
    booking_counts = _db.get_slot_booking_counts(date_val)
    # Build full slot list: no record = available (admin closes slots manually)
    all_slots = []
    available_map = {e["time_slot"]: e["is_available"] for e in entries}
    for ts in _db.TIME_SLOTS:
        is_open = available_map.get(ts, True)  # Default: available
        booked = booking_counts.get(ts, 0)
        is_av = is_open and booked < SLOT_CAPACITY
        all_slots.append({
            "time_slot": ts,
            "is_available": is_av,
            "is_open": is_open,
            "booked_count": booked,
            "capacity": SLOT_CAPACITY,
        })
    return jsonify({"ok": True, "date": date_str, "slots": all_slots})


@bp.route("/api/availability/month")
def api_availability_month():
    """Get availability + booking counts for next 30 days (for admin calendar management)."""
    today = date.today()
    end = today + timedelta(days=29)
    range_data = _db.get_availability_range(today, end)
    booking_counts = _db.get_booking_counts_range(today, end)
    return jsonify({
        "ok": True,
        "start": today.isoformat(),
        "end": end.isoformat(),
        "data": range_data,
        "bookings": booking_counts,
    })


@bp.route("/api/availability", methods=["POST"])
@role_required('admin')
def api_set_availability():
    """Batch set availability for a date. Admin only.
    Body: {"date": "YYYY-MM-DD", "slots": [{"time_slot": "09:00", "is_available": true}, ...]}
    """
    data = request.get_json(force=True)
    date_str = data.get("date", "")
    slots = data.get("slots", [])
    if not date_str:
        return jsonify({"ok": False, "error": "请指定日期"}), 400
    try:
        date_val = date.fromisoformat(date_str)
    except ValueError:
        return jsonify({"ok": False, "error": "日期格式错误"}), 400
    _db.batch_set_availability(date_val, slots)
    return jsonify({"ok": True})
@bp.route("/api/booking/<int:booking_id>", methods=["DELETE"])
@admin_required
def api_delete_booking(booking_id):
    try:
        _db.delete_booking(booking_id)
        return jsonify({"ok": True})
    except Exception as e:
        return jsonify({"ok": False, "error": str(e)}), 500
PROMPT_VERSIONS_DIR = BASE_DIR / "prompts" / "versions"
PROMPT_VERSIONS_DIR.mkdir(parents=True, exist_ok=True)


def _get_next_prompt_version() -> int:
    """Scan prompts/versions/ and return next version number."""
    existing = sorted(PROMPT_VERSIONS_DIR.glob("ai_interpreter_v*.md"))
    if not existing:
        return 1
    nums = []
    for f in existing:
        try:
            n = int(f.stem.split("_v")[1])
            nums.append(n)
        except (ValueError, IndexError):
            pass
    return max(nums) + 1 if nums else 1


@bp.route("/prompt-lab")
@role_required('admin')
def prompt_lab_page():
    """Prompt iteration test bench."""
    return render_template("prompt_lab.html")


@bp.route("/internal")
@role_required('admin', 'advisor')
def internal_download_page():
    """内部下载页：有效评价 qualification + 协议下载（生产版，不含 prompt 迭代工具）。"""
    return render_template("internal.html")
@bp.route("/api/prompt-lab/run", methods=["POST"])
@role_required('admin')
def prompt_lab_run():
    """Run AI interpretation using current prompt + report_data.json."""
    import urllib.request as _ureq
    import time as _time

    data = request.get_json(force=True)
    user_message = (data.get("message") or "请给出这份 Y4 报告的完整解读").strip()

    # 1) Read prompt
    prompt_path = BASE_DIR / "prompts" / "ai_interpreter.md"
    system_prompt = prompt_path.read_text(encoding="utf-8") if prompt_path.exists() else "你是测评解读助手。"

    # 2) Read report data
    report_path = DATA_DIR / "report_data.json"
    if not report_path.exists():
        return jsonify({"ok": False, "error": "没有测试数据 (report_data.json 不存在)"}), 400

    report_data = json.loads(report_path.read_text(encoding="utf-8"))
    schema_items = report_data.get("schema_124", [])
    student = report_data.get("student", {})
    context = _build_y4_interpret_context(schema_items, student)

    # 3) Build messages
    messages = [
        {"role": "system", "content": system_prompt + "\n\n以下是 Y4 有效评价数据：\n" + context},
        {"role": "user", "content": user_message},
    ]

    # 4) Call DashScope
    dashscope_key = os.environ.get("DASHSCOPE_API_KEY", extract.DEFAULT_DASHSCOPE_KEY).strip()
    url = "https://dashscope.aliyuncs.com/compatible-mode/v1/chat/completions"
    payload = json.dumps({
        "model": os.environ.get("AI_TEXT_MODEL", "qwen-plus"),
        "messages": messages,
        "temperature": float(os.environ.get("AI_TEMPERATURE", "0.3")),
        "max_tokens": 8192,
    }).encode("utf-8")
    req = _ureq.Request(
        url, data=payload,
        headers={"Authorization": f"Bearer {dashscope_key}",
                 "Content-Type": "application/json"},
        method="POST",
    )
    t0 = _time.time()
    try:
        with _ureq.urlopen(req, timeout=120) as resp:
            result = json.loads(resp.read().decode("utf-8"))
        reply = _finalize_y4_interpretation(
            result["choices"][0]["message"]["content"], dashscope_key)
        tokens_used = result.get("usage", {}).get("total_tokens", 0)
        elapsed_ms = int((_time.time() - t0) * 1000)
        return jsonify({"ok": True, "reply": reply, "tokens": tokens_used, "time_ms": elapsed_ms})
    except Exception as exc:
        return jsonify({"ok": False, "error": f"AI 调用失败: {exc}"}), 500


@bp.route("/api/prompt-lab/iterate", methods=["POST"])
@role_required('admin')
def prompt_lab_iterate():
    """Auto-modify prompt based on user feedback, then return new prompt."""
    import urllib.request as _ureq

    data = request.get_json(force=True)
    feedback = (data.get("feedback") or "").strip()
    rating = data.get("rating", 3)
    last_output = (data.get("last_output") or "").strip()

    if not feedback:
        return jsonify({"ok": False, "error": "反馈不能为空"}), 400

    # 1) Read current prompt
    prompt_path = BASE_DIR / "prompts" / "ai_interpreter.md"
    current_prompt = prompt_path.read_text(encoding="utf-8") if prompt_path.exists() else ""

    # 2) Save current version
    version_num = _get_next_prompt_version()
    old_version_path = PROMPT_VERSIONS_DIR / f"ai_interpreter_v{version_num}.md"
    old_version_path.write_text(current_prompt, encoding="utf-8")

    # 3) Build meta-prompt for AI to improve the prompt
    meta_prompt = f"""你是一个 prompt 优化专家。你需要根据用户的反馈，修改 Y4 测评解读的 system prompt。

以下是当前用于 Y4 测评解读的 system prompt：
---
{current_prompt}
---

以下是用户对这个 prompt 生成输出的反馈：
评分：{rating}/5
反馈：{feedback}

上次 AI 的输出（供参考）：
---
{last_output[:3000]}
---

请根据用户反馈，修改上面的 system prompt。
要求：
- 只修改需要改进的部分，保留好的部分
- 输出完整的修改后的 prompt（不是 diff，不是解释，直接输出 prompt 全文）
- 保持 Y4 四维框架（心力/精力/学习力/生涯力）作为唯一语言体系
- 保持指标编号引用规范
- 不要加任何前缀说明或后缀解释"""

    messages = [{"role": "user", "content": meta_prompt}]

    # 4) Call DashScope to get improved prompt
    dashscope_key = os.environ.get("DASHSCOPE_API_KEY", extract.DEFAULT_DASHSCOPE_KEY).strip()
    url = "https://dashscope.aliyuncs.com/compatible-mode/v1/chat/completions"
    payload = json.dumps({
        "model": os.environ.get("AI_TEXT_MODEL", "qwen-plus"),
        "messages": messages,
        "temperature": 0.3,
        "max_tokens": 8192,
    }).encode("utf-8")
    req = _ureq.Request(
        url, data=payload,
        headers={"Authorization": f"Bearer {dashscope_key}",
                 "Content-Type": "application/json"},
        method="POST",
    )
    try:
        with _ureq.urlopen(req, timeout=120) as resp:
            result = json.loads(resp.read().decode("utf-8"))
        new_prompt = result["choices"][0]["message"]["content"].strip()

        # 5) Write new prompt
        prompt_path.write_text(new_prompt, encoding="utf-8")

        # 6) Simple diff: find first and last differing lines
        old_lines = current_prompt.split("\n")
        new_lines = new_prompt.split("\n")
        diff_parts = []
        max_compare = min(len(old_lines), len(new_lines))
        first_diff = None
        last_diff = None
        for i in range(max_compare):
            if old_lines[i] != new_lines[i]:
                if first_diff is None:
                    first_diff = i
                last_diff = i
        if first_diff is not None:
            start = max(0, first_diff - 2)
            end = min(max_compare, last_diff + 3)
            diff_parts.append(f"改动区域（第 {start+1}-{end} 行附近）：")
            for i in range(start, end):
                marker = "→" if old_lines[i] != new_lines[i] else " "
                if old_lines[i] != new_lines[i]:
                    diff_parts.append(f"  {marker} 旧: {old_lines[i][:80]}")
                    diff_parts.append(f"  {marker} 新: {new_lines[i][:80]}")
        if len(new_lines) != len(old_lines):
            diff_parts.append(f"\n行数变化：{len(old_lines)} → {len(new_lines)}")
        diff_summary = "\n".join(diff_parts) if diff_parts else "无明显差异"

        return jsonify({
            "ok": True,
            "old_prompt": current_prompt,
            "new_prompt": new_prompt,
            "diff": diff_summary,
            "version": version_num,
        })
    except Exception as exc:
        return jsonify({"ok": False, "error": f"Prompt 迭代失败: {exc}"}), 500


@bp.route("/api/prompt-lab/save", methods=["POST"])
@role_required('admin')
def prompt_lab_save():
    """Manually save edited prompt."""
    data = request.get_json(force=True)
    prompt_text = data.get("prompt", "").strip()
    if not prompt_text:
        return jsonify({"ok": False, "error": "Prompt 不能为空"}), 400

    # Save old version first
    prompt_path = BASE_DIR / "prompts" / "ai_interpreter.md"
    if prompt_path.exists():
        version_num = _get_next_prompt_version()
        old_version_path = PROMPT_VERSIONS_DIR / f"ai_interpreter_v{version_num}.md"
        old_version_path.write_text(prompt_path.read_text(encoding="utf-8"), encoding="utf-8")

    prompt_path.write_text(prompt_text, encoding="utf-8")
    return jsonify({"ok": True})


@bp.route("/api/prompt-lab/prompt", methods=["GET"])
@role_required('admin')
def prompt_lab_get_prompt():
    """Get current prompt text."""
    prompt_path = BASE_DIR / "prompts" / "ai_interpreter.md"
    content = prompt_path.read_text(encoding="utf-8") if prompt_path.exists() else ""
    return jsonify({"ok": True, "prompt": content})
@bp.route("/api/prompt-lab/evaluation")
@role_required('admin')
def prompt_lab_evaluation():
    """返回 report_data.json 的有效评价视图数据点。"""
    report_path = DATA_DIR / "report_data.json"
    if not report_path.exists():
        return jsonify({"ok": False, "error": "没有测试数据 (report_data.json 不存在)"}), 400
    report_data = json.loads(report_path.read_text(encoding="utf-8"))
    schema_items = report_data.get("schema_124", [])
    student = report_data.get("student", {}) or {}
    items, summary, mapping_done, per_code_opts = _build_qualification_view(schema_items)
    return jsonify({"ok": True, "items": items, "student": student, "summary": summary,
                    "mapping_done": mapping_done, "mapping_total": len(items),
                    "data_bias_options": _eval_rules.DATA_BIAS_OPTIONS,
                    "per_code_eval_options": per_code_opts})
@bp.route("/api/prompt-lab/evaluation/download", methods=["POST"])
@role_required('admin')
def prompt_lab_evaluation_download():
    """下载 AI→AI 传递协议 Markdown（Phase 1）。

    入参: {items: [{code,label,raw_value,eval_value,dimension,unit,eval_source,note}],
           include_raw: bool, interpretation: str(可选)}
    """
    import io as _io
    data = request.get_json(force=True)
    items = data.get("items", [])
    include_raw = bool(data.get("include_raw", False))
    interpretation = (data.get("interpretation") or "").strip()

    # 取 student 信息
    report_path = DATA_DIR / "report_data.json"
    student = {}
    if report_path.exists():
        student = json.loads(report_path.read_text(encoding="utf-8")).get("student", {}) or {}
    student_name = student.get("name") or "测试"

    md_text = _build_phase1_md_text(items, include_raw, interpretation, student, student_name)
    buf = _io.BytesIO(md_text.encode("utf-8"))
    safe_name = student_name.replace(" ", "").replace("/", "_")
    return send_file(buf, mimetype="text/markdown", as_attachment=True,
                     download_name=f"Y4报告_{safe_name}.md")
@bp.route("/api/prompt-lab/e4-mapping", methods=["GET"])
@role_required('admin', 'advisor')
def prompt_lab_e4_mapping_get():
    """返回用户指认的 Y4→E4 映射及框架常量。"""
    return jsonify({
        "ok": True,
        "mapping": _db.get_e4_mapping(),
        "dims": _eval_rules.E4_LABELS,
        "subcats": _eval_rules.E4_SUBCATS,
    })


@bp.route("/api/prompt-lab/e4-mapping", methods=["POST"])
@role_required('admin')
def prompt_lab_e4_mapping_save():
    """保存映射指认。body: {entries: [{code, e4_dims: [{e4_dim, e4_sub}, ...]}]}。

    e4_dims 为空列表表示清除该 code 的全部映射。
    e4_sub 可空（子类别为可选）。
    """
    data = request.get_json(force=True)
    entries = data.get("entries") or []
    cleaned = []
    for e in entries:
        code = str(e.get("code", "")).strip()
        if not code:
            continue
        dims = e.get("e4_dims") or []
        clean_dims = []
        for d in dims:
            dim = (d.get("e4_dim") or "").strip()
            if not dim:
                continue
            if dim not in _eval_rules.E4_DIMS:
                return jsonify({"ok": False, "error": f"未知 E4 维度: {dim}"}), 400
            sub = (d.get("e4_sub") or "").strip()
            # 子类别可选：如果填了，校验是否在预设列表里；不填也允许
            if sub and sub not in _eval_rules.E4_SUBCATS.get(dim, []):
                # 允许用户新建子类别（在 E4 框架线内），不报错
                pass
            clean_dims.append({"e4_dim": dim, "e4_sub": sub})
        cleaned.append({"code": code, "e4_dims": clean_dims})
    mapping = _db.save_e4_mapping(cleaned)
    return jsonify({"ok": True, "mapping": mapping})


@bp.route("/api/prompt-lab/e4/step1", methods=["POST"])
@role_required('admin')
def prompt_lab_e4_step1():
    """E4 工作流 Step1：数据整编（纯 Python，无 AI）。

    按 E4 框架 source of truth（维度→评估问题→指标）分组。
    框架未引用的指标归入 OTHER。
    """
    data = request.get_json(force=True)
    items = data.get("items", [])
    include_raw = bool(data.get("include_raw", False))

    vgroups, other_items, _dims_by_code = _eval_rules.evaluate_question_verdicts(items)

    groups: List[Dict[str, Any]] = []
    group_counts: Dict[str, int] = {d: 0 for d in _eval_rules.E4_DIMS}
    for g in vgroups:
        lines = [_format_e4_line(it, include_raw) for it in g["items"]]
        groups.append({"dim": g["dim"], "q": g["q"], "lines": lines, "verdict": g.get("verdict")})
        group_counts[g["dim"]] += len(g["items"])
    other_lines = [_format_e4_line(it, include_raw) for it in other_items]

    # 四类名单预览（干预/潜能 × 强/弱）
    list_placement = _determine_list_placement(items)

    return jsonify({
        "ok": True,
        "groups": groups,
        "other_lines": other_lines,
        "group_counts": group_counts,
        "unmapped_count": len(other_lines),
        "list_placement": list_placement,
    })


@bp.route("/api/prompt-lab/e4/step2", methods=["POST"])
@role_required('admin')
def prompt_lab_e4_step2():
    """E4 工作流 Step2：基于框架问题的 contextual 分析。

    数据按 E4 框架 source of truth 的「维度→评估问题→指标」组织。
    AI 的任务是逐一回答框架问题，并输出三段式综合（核心问题/可能失真·先观察/跟进线索与切入点）。
    """
    data = request.get_json(force=True)
    groups = data.get("groups") or []
    other_lines = data.get("other_lines") or []
    model = os.environ.get("E4_STEP2_MODEL", "qwen-turbo")
    timeout = int(os.environ.get("E4_STEP_TIMEOUT", "180"))

    # 按维度+问题拼装数据（附 Python 规则的问题级判定）
    all_lines: List[str] = []
    current_dim = None
    for g in groups:
        dim = g.get("dim", "")
        if dim != current_dim:
            all_lines.append("")
            all_lines.append(f"## {_eval_rules.E4_LABELS.get(dim, dim)}")
            current_dim = dim
        q = g.get("q")
        if q:
            all_lines.append(f"问题：{q}")
        v = g.get("verdict")
        if v:
            cn = _eval_rules.verdict_label_cn(v.get("state"), v.get("items") or [])
            la = v.get("list_assignment") or ""
            la_seg = f" | 名单: {la}" if la else ""
            all_lines.append(f"问题判定（Python 规则，须以此为准绳）：{cn}{la_seg} | {v.get('summary', '')}")
        all_lines.extend(g.get("lines") or [])
        all_lines.append("")
    if other_lines:
        all_lines.append("## OTHER_VARIABLES（框架未引用）")
        all_lines.extend(other_lines)
        all_lines.append("")

    if not any(g.get("lines") for g in groups):
        return jsonify({"ok": True, "analyses": {d: "（无数据）" for d in _eval_rules.E4_DIMS},
                        "stats": [], "failed": [], "model": model, "full_analysis": ""})

    system = _e4_step2_system()
    user = ("【E4 框架数据（按维度和评估问题组织；问题下列出该问题涉及的指标数据）】\n"
            + "\n".join(all_lines))

    try:
        r = _dashscope_chat(
            [{"role": "system", "content": system},
             {"role": "user", "content": user}],
            model=model, timeout=timeout, max_tokens=4000)
        full_analysis = _sanitize_e4_step2(r["content"])
        analyses: Dict[str, str] = {d: "" for d in _eval_rules.E4_DIMS}
        analyses["_cross"] = _e4_step2_tail(full_analysis)

        stats = [{"dim": "ALL", "tokens": r["tokens"], "time_ms": r["time_ms"], "model": model}]
        return jsonify({"ok": True, "analyses": analyses,
                        "stats": stats, "failed": [], "model": model,
                        "full_analysis": full_analysis})
    except Exception as exc:
        return jsonify({"ok": False, "analyses": {}, "stats": [], "failed": ["ALL"],
                        "model": model, "error": str(exc)})


@bp.route("/api/prompt-lab/e4/step3", methods=["POST"])
@role_required('admin')
def prompt_lab_e4_step3():
    """E4 工作流 Step3：综合成文。

    AI 只写「总览」段；最终协议由 Python 确定性拼装（AI 不碰数据行，保证 Y4 原名）。
    不落库，直接返回 content_md。
    """
    data = request.get_json(force=True)
    items = data.get("items", [])
    include_raw = bool(data.get("include_raw", False))
    analyses = data.get("analyses") or {}
    full_analysis = (data.get("full_analysis") or "").strip()
    interpretation = (data.get("interpretation") or "").strip()
    model = os.environ.get("E4_STEP3_MODEL", "qwen-plus")
    timeout = int(os.environ.get("E4_STEP_TIMEOUT", "180"))

    # 规则兜底名单判定（AI 只写理由，不得改判）
    list_placement = _determine_list_placement(items, mapping=None)

    # 1) AI 写总览（名单归属以 Python 规则为准）
    system = _e4_step3_system(list_placement.get("category")
                              or list_placement.get("placement") or "（未判定）")
    user = (f"【综合分析草稿】\n{full_analysis or analyses}\n\n"
            f"【Y4 原始 AI 解读（可选参考，非客观事实）】\n{interpretation or '（无）'}")
    try:
        r = _dashscope_chat(
            [{"role": "system", "content": system},
             {"role": "user", "content": user}],
            model=model, timeout=timeout, max_tokens=2500)
        overview_md = r["content"]
        stats = {"tokens": r["tokens"], "time_ms": r["time_ms"], "model": model}
    except Exception as exc:
        overview_md = f"（总览生成失败：{exc}）"
        stats = {"tokens": 0, "time_ms": 0, "model": model, "error": str(exc)}

    # 2) Python 确定性拼装最终协议
    student_name = _get_student_name_from_items(items)
    content_md = _assemble_e4_protocol(items, include_raw, analyses, overview_md, full_analysis, interpretation)
    return jsonify({"ok": True, "content_md": content_md, "stats": stats,
                    "student_name": student_name, "list_placement": list_placement})


@bp.route("/api/prompt-lab/y4-export/json", methods=["POST"])
@role_required('admin')
def prompt_lab_y4_export_json():
    """Generalized Y4 JSON 导出（非 E4）：涵盖全部数据点 + AI 解读。

    不含 code、不含 E4 框架分组；按 Y4 四维组织。
    """
    data = request.get_json(force=True)
    items = data.get("items", [])
    interpretation = (data.get("interpretation") or "").strip()

    report_path = DATA_DIR / "report_data.json"
    student: Dict[str, Any] = {}
    if report_path.exists():
        student = json.loads(report_path.read_text(encoding="utf-8")).get("student", {}) or {}
    student_name = student.get("name") or "测试"

    payload = _build_y4_payload(items, interpretation, student, student_name)
    return jsonify({"ok": True, "json": payload,
                    "filename": f"Y4数据_{student_name.replace(' ', '').replace('/', '_')}.json"})
