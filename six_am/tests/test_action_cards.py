# -*- coding: utf-8 -*-
"""Task 3：行动卡视图模型与固定行动模板测试。

覆盖（实施计划 Task 3 Step 1）：
- test_hypothesis_card_uses_y4_facts_and_confirmation_language
- test_confirmed_card_uses_current_evidence_language
- test_a3_guidance_says_restore_capacity_before_reassessing_a4_a5
- test_a6_card_uses_fit_label_not_score
- test_action_overrides_replace_generated_actions
- test_v1_profile_with_score_degrades_without_error
另加 X.7 场景：A1 必须显示 priority_label=长期目标 与 deferred_reason。

运行：python3.11 -m unittest six_am.tests.test_action_cards -v
"""

import json
import os
import sys
import unittest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))

from six_am import engine  # noqa: E402
from six_am import action_guidance, presentation  # noqa: E402


def item(label, value):
    return {"code": "", "label": label, "raw_value": value, "eval_value": value,
            "eval_source": "规则推导", "dimension": "", "level_code": None,
            "level_label": None, "data_bias": "正常"}


def order(value, pos, total=15):
    return {"code": "%03d" % (109 + pos), "label": "职业价值观排序%d" % pos,
            "raw_value": value, "eval_value": value, "eval_source": "原始排序",
            "dimension": "生涯力", "data_bias": "正常"}


def run(items, evidence=None, answers=None):
    return engine.run(items, student={"name": "测试学生", "grade": "高二"},
                      evidence=evidence, answers=answers)


def rebellion_items():
    """v0.2 §10 例 1：逆反型（A1 Blocked / A1-S2 / Hypothesis）。"""
    return [
        item("自驱力-自主性", "偏低"),
        order("独立自主", 2),
        item("依恋关系-沟通-母亲", "偏低"),
        item("人格-宜人性", "不低"),
        item("职业兴趣-研究型", "高"),
        item("职业兴趣-现实型", "不低"),
        item("职业兴趣-艺术型", "不低"),
        item("职业兴趣-社会型", "偏低"),
        item("职业兴趣-事业型", "不低"),
        item("职业兴趣-常规型", "偏低"),
        item("学习动机-深层动机", "不低"),
        item("体质健康-睡眠评级", "良好"),
        item("学习动机-自我效能感", "不低"),
    ]


def overload_items():
    """v0.2 §10 例 2：Capacity Overload（A3 Blocked / A3-S1）。"""
    return [
        item("体质健康-睡眠评级", "差"),
        item("体质健康-运动评级", "一般"),
        item("情绪稳定性-焦虑安详", "需关注"),
        item("情绪稳定性总分", "偏低"),
        item("学习动机-自我效能感", "偏低"),
        item("学习方法与策略-学习自我调节", "偏低"),
        item("人格-责任心", "偏低"),
    ]


def cards_by_dim(vm):
    return {c["dim"]: c for c in vm["cards"]}


CARD_FIELDS = ["dim", "title", "state", "state_label", "evidence_state", "evidence_label",
               "headline", "y4_facts", "hypotheses", "confirmed_evidence",
               "next_actions", "avoid_actions", "questions", "audit",
               "priority_rank", "priority_reason", "priority_label",
               "deferred", "deferred_reason", "manually_edited"]


class HypothesisCard(unittest.TestCase):
    """仅线索（Hypothesis）卡：Y4 事实 + “可能指向…需要进一步确认”。"""

    def setUp(self):
        self.profile = run(rebellion_items())
        self.vm = presentation.build_view_model(self.profile)
        self.cards = cards_by_dim(self.vm)

    def test_hypothesis_card_uses_y4_facts_and_confirmation_language(self):
        a1 = self.cards["A1"]
        for field in CARD_FIELDS:
            self.assertIn(field, a1, field)
        self.assertEqual(a1["state"], "Blocked")
        self.assertEqual(a1["state_label"], "阻塞 Blocked")
        self.assertEqual(a1["evidence_state"], "Hypothesis")
        self.assertEqual(a1["evidence_label"], "仅线索 Hypothesis")
        # 固定模板：Y4 显示：{Y4事实}；可能指向：{规则标题}，需要进一步确认。
        self.assertTrue(a1["headline"].startswith("Y4 显示："), a1["headline"])
        self.assertIn("可能指向：", a1["headline"])
        self.assertTrue(a1["headline"].endswith("，需要进一步确认。"), a1["headline"])
        # Y4 原词：指标名 + 值
        self.assertIn("自驱力-自主性", a1["headline"])
        self.assertIn("偏低", a1["headline"])
        # 规则标题进入 headline 与 hypotheses
        self.assertIn("高叛逆风险", a1["headline"])
        self.assertTrue(any("自驱力-自主性" in f and "偏低" in f for f in a1["y4_facts"]))
        self.assertTrue(any("高叛逆风险" in h for h in a1["hypotheses"]))
        # 不得把线索写成结论
        self.assertNotIn("当前判断为", a1["headline"])

    def test_secure_card_uses_secure_template(self):
        a3 = self.cards["A3"]
        self.assertEqual(a3["state"], "Secure")
        self.assertEqual(a3["headline"], "现有 Y4 事实未提示该维度构成启动障碍。")


class ConfirmedCard(unittest.TestCase):
    """已证实（Confirmed）卡：根据目前证据，当前判断为…；依据包括…。"""

    def setUp(self):
        answers = [{"q": "Q1.3", "result": "confirm"}]
        self.profile = run(rebellion_items(), answers=answers)
        self.vm = presentation.build_view_model(self.profile)
        self.cards = cards_by_dim(self.vm)

    def test_confirmed_card_uses_current_evidence_language(self):
        a1 = self.cards["A1"]
        self.assertEqual(a1["evidence_state"], "Confirmed")
        self.assertEqual(a1["evidence_label"], "已证实 Confirmed")
        # 固定模板：根据目前证据，当前判断为：{状态中文}{亚型}；依据包括：{证据原文}。
        self.assertTrue(a1["headline"].startswith("根据目前证据，当前判断为："), a1["headline"])
        self.assertIn("阻塞", a1["headline"])
        self.assertIn("逆反", a1["headline"])
        self.assertIn("依据包括：", a1["headline"])
        self.assertTrue(a1["confirmed_evidence"])
        # 依据含规则原文/标题，不用“可能”口吻
        self.assertNotIn("需要进一步确认", a1["headline"])


class A3Guidance(unittest.TestCase):
    """A3 过载：先恢复容量，再复评 A4/A5（简明版 §4 处理原则 / §5 联动）。"""

    def setUp(self):
        evidence = {"ft": {"schedule": {"load_level": "高", "free_hours_per_week": 2}}}
        self.profile = run(overload_items(), evidence=evidence)
        self.vm = presentation.build_view_model(self.profile)
        self.cards = cards_by_dim(self.vm)

    def test_a3_guidance_says_restore_capacity_before_reassessing_a4_a5(self):
        a3 = self.profile["dimensions"]["A3"]
        self.assertEqual(a3["subtype"], "A3-S1")
        g = action_guidance.guidance_for("A3", a3, self.profile["support_plan"])
        self.assertTrue(any("睡眠" in a or "恢复" in a for a in g["next_actions"]),
                        g["next_actions"])
        self.assertTrue(any("先处理容量问题" in a and "复评 A4/A5" in a
                            for a in g["next_actions"]), g["next_actions"])
        card = self.cards["A3"]
        self.assertTrue(any("复评 A4/A5" in a for a in card["next_actions"]))


class A6Card(unittest.TestCase):
    """A6 支持选择器：只显示定性适配标签，不显示数值分。"""

    def setUp(self):
        self.profile = run(rebellion_items())
        self.vm = presentation.build_view_model(self.profile)
        self.a6 = cards_by_dim(self.vm)["A6"]

    def test_a6_card_uses_fit_label_not_score(self):
        a6 = self.a6
        self.assertEqual(a6["state_label"], "支持选择器")
        # 固定模板：根据前五维结论，优先尝试：{支持方向}；需要避免：{避免方向}。
        self.assertTrue(a6["headline"].startswith("根据前五维结论，优先尝试："), a6["headline"])
        self.assertIn("同伴协作", a6["headline"])
        self.assertIn("需要避免：", a6["headline"])
        self.assertIn("慎用", a6["headline"])
        # next/avoid 来自 support_plan 的 fit_label / cautions
        self.assertTrue(any("优先" in a and "同伴协作" in a for a in a6["next_actions"]),
                        a6["next_actions"])
        self.assertTrue(any("慎用" in a for a in a6["avoid_actions"]), a6["avoid_actions"])
        # 全卡不得出现数值分或预测有效性
        blob = json.dumps(a6, ensure_ascii=False)
        self.assertNotIn('"score"', blob)
        self.assertNotIn("predicted_effectiveness", blob)


class ActionOverrides(unittest.TestCase):
    """人工行动修订替换生成的行动；空 override 回退到生成结果。"""

    def setUp(self):
        self.profile = run(rebellion_items())

    def test_action_overrides_replace_generated_actions(self):
        self.profile["action_overrides"] = {
            "A1": {"next_actions": ["和家长协商一个可修改的目标"],
                   "avoid_actions": ["本周不追加任何监督"]},
        }
        vm = presentation.build_view_model(self.profile)
        cards = cards_by_dim(vm)
        a1 = cards["A1"]
        self.assertEqual(a1["next_actions"], ["和家长协商一个可修改的目标"])
        self.assertEqual(a1["avoid_actions"], ["本周不追加任何监督"])
        self.assertTrue(a1["manually_edited"])
        # 未修订维度仍用生成结果
        self.assertFalse(cards["A2"]["manually_edited"])
        self.assertTrue(cards["A2"]["next_actions"])

    def test_view_model_applies_action_overrides(self):
        """视图模型层面：override 完整替换生成行动并标记 manually_edited。"""
        override = {"next_actions": ["先睡够再谈目标"],
                    "avoid_actions": ["不要同时改三件事"]}
        self.profile["action_overrides"] = {"A3": override}
        vm = presentation.build_view_model(self.profile)
        a3 = cards_by_dim(vm)["A3"]
        self.assertEqual(a3["next_actions"], override["next_actions"])
        self.assertEqual(a3["avoid_actions"], override["avoid_actions"])
        self.assertTrue(a3["manually_edited"])
        # 摘要不受修订影响
        self.assertIn("summary", vm)

    def test_empty_override_falls_back_to_generated_actions(self):
        generated = presentation.build_view_model(self.profile)
        gen_a1 = cards_by_dim(generated)["A1"]
        self.profile["action_overrides"] = {"A1": {"next_actions": [], "avoid_actions": []}}
        vm = presentation.build_view_model(self.profile)
        a1 = cards_by_dim(vm)["A1"]
        self.assertEqual(a1["next_actions"], gen_a1["next_actions"])
        self.assertEqual(a1["avoid_actions"], gen_a1["avoid_actions"])
        self.assertFalse(a1["manually_edited"])


class X7DeferredCard(unittest.TestCase):
    """X.7：A1 不进入常规障碍排序，显示长期目标。"""

    def setUp(self):
        evidence = {"parent": {"cannot_relinquish": True}}
        self.profile = run(rebellion_items(), evidence=evidence)
        self.vm = presentation.build_view_model(self.profile)
        self.a1 = cards_by_dim(self.vm)["A1"]

    def test_x7_a1_shows_long_term_goal_label_and_deferred_reason(self):
        a1 = self.a1
        self.assertTrue(a1["deferred"])
        self.assertEqual(a1["priority_label"], "长期目标")
        self.assertTrue(a1["deferred_reason"])
        self.assertIn("家长", a1["deferred_reason"])
        self.assertIsNone(a1["priority_rank"])


class EvidencePathNoLeak(unittest.TestCase):
    """仅线索卡不得在 confirmed_evidence 暴露内部 {"evidence": path} 条目。"""

    def test_hypothesis_card_never_exposes_internal_evidence_paths(self):
        items = [
            item("职业兴趣-研究型", "高"),      # holland_peak_exists
            item("人格-宜人性", "不低"),
            item("自驱力-自主性", "不低"),
        ]
        # task.mismatch=True → R2.2 任务错位命中（Hypothesis），其证据含
        # {"evidence": "task.mismatch", "value": True} 内部路径条目
        profile = run(items, evidence={"task": {"mismatch": True}})
        vm = presentation.build_view_model(profile)
        a2 = cards_by_dim(vm)["A2"]
        self.assertEqual(a2["evidence_state"], "Hypothesis")
        self.assertEqual(a2["confirmed_evidence"], [])
        # 主卡字段（除 audit 外）序列化后不得出现内部证据路径或布尔值
        main = {k: v for k, v in a2.items() if k != "audit"}
        blob = json.dumps(main, ensure_ascii=False)
        self.assertNotIn("task.mismatch", blob)
        # 原始证据条目仍保留在 audit 规则链中
        audit_blob = json.dumps(a2["audit"], ensure_ascii=False)
        self.assertIn("task.mismatch", audit_blob)

    def test_confirmed_card_keeps_rule_titles_without_evidence_paths(self):
        items = [
            item("职业兴趣-研究型", "高"),
            item("人格-宜人性", "不低"),
            item("自驱力-自主性", "不低"),
        ]
        answers = [{"q": "Q2.1", "result": "confirm"}]
        profile = run(items, evidence={"task": {"mismatch": True}}, answers=answers)
        vm = presentation.build_view_model(profile)
        a2 = cards_by_dim(vm)["A2"]
        self.assertEqual(a2["evidence_state"], "Confirmed")
        self.assertTrue(a2["confirmed_evidence"])
        blob = json.dumps(a2["confirmed_evidence"], ensure_ascii=False)
        self.assertNotIn("task.mismatch", blob)
        self.assertNotIn("True", blob)


class RefutedSurfacing(unittest.TestCase):
    """被追问否定的规则在视图模型中可见，可渲染「被排除 Refuted」。"""

    def setUp(self):
        answers = [{"q": "Q1.3", "result": "deny"}]   # 否定 R1.3
        self.profile = run(rebellion_items(), answers=answers)
        self.assertEqual(self.profile["refuted"], {"R1.3": "Q1.3"})
        self.vm = presentation.build_view_model(self.profile)

    def test_refuted_rule_surfaced_in_view_model(self):
        blob = json.dumps(self.vm, ensure_ascii=False)
        self.assertIn("被排除 Refuted", blob)
        # 摘要级：规则 ID、标题、否定来源追问
        refuted = self.vm["summary"]["refuted"]
        self.assertEqual(len(refuted), 1)
        entry = refuted[0]
        self.assertEqual(entry["rule"], "R1.3")
        self.assertEqual(entry["title"], "高叛逆风险（双条件）")
        self.assertEqual(entry["refuted_by"], "Q1.3")
        self.assertEqual(entry["dim"], "A1")
        self.assertEqual(entry["label"], "被排除 Refuted")
        # 卡级：A1 audit 含被排除规则，其他维度不含
        cards = cards_by_dim(self.vm)
        a1_refuted = cards["A1"]["audit"]["refuted_rules"]
        self.assertEqual([r["rule"] for r in a1_refuted], ["R1.3"])
        self.assertEqual(a1_refuted[0]["label"], "被排除 Refuted")
        for dim in ("A2", "A3", "A4", "A5", "A6"):
            self.assertEqual(cards[dim]["audit"]["refuted_rules"], [], dim)

    def test_no_refuted_means_empty_everywhere(self):
        vm = presentation.build_view_model(run(rebellion_items()))
        self.assertEqual(vm["summary"]["refuted"], [])
        blob = json.dumps(vm, ensure_ascii=False)
        self.assertNotIn("被排除 Refuted", blob)
        for card in vm["cards"]:
            self.assertEqual(card["audit"]["refuted_rules"], [])


class V1Degrade(unittest.TestCase):
    """迁移窗口内的 v1 profile：忽略旧 score，缺 priority_rank 时按固定顺序降级。"""

    def test_v1_profile_with_score_degrades_without_error(self):
        v1 = {
            "schema": "six-am-profile-v1",
            "student": {"name": "旧版学生", "grade": "高一"},
            "report_id": 7,
            "dimensions": {
                "A1": {"state": "Blocked", "subtype": "A1-S2",
                       "subtype_name": "Rebellion 逆反型（高叛逆风险）",
                       "confidence": "Hypothesis",
                       "score": {"total": 12.3, "base": 3.0},
                       "triggered_rules": [
                           {"id": "R1.3", "title": "高叛逆风险（双条件）",
                            "evidence": [{"point": "自驱力-自主性", "eval_value": "偏低",
                                          "source": "规则推导"}],
                            "effect": {"state": "Blocked"}, "confidence": "Hypothesis",
                            "contributing": True}],
                       "cross_mods": [], "flags": [], "notes": []},
                "A2": {"state": "Watch", "subtype": None, "subtype_name": None,
                       "confidence": "Hypothesis", "score": {"total": 4.0},
                       "triggered_rules": [], "cross_mods": [], "flags": [], "notes": []},
                "A3": {"state": "Secure", "score": {"total": 0.0}},
                "A4": {"state": "Secure", "score": {"total": 0.0}},
                "A5": {"state": "Secure", "score": {"total": 0.0}},
                "A6": {"state": None, "score": {"total": 0.0}},
            },
            # 无 priority_order / priority_rank / support_plan / deferred / questions
            # v1 旧障碍条目带 score/high_priority（旧 scoring.compute 产出）
            "obstacles": [{"dim": "A1", "score": 12.3, "high_priority": True,
                           "sequence": 1, "sequence_note": ""}],
        }
        vm = presentation.build_view_model(v1)
        self.assertEqual([c["dim"] for c in vm["cards"]],
                         ["A1", "A2", "A3", "A4", "A5", "A6"])
        for card in vm["cards"]:
            for field in CARD_FIELDS:
                self.assertIn(field, card, field)
            self.assertIsNone(card["priority_rank"])
            self.assertFalse(card["deferred"])
            self.assertFalse(card["manually_edited"])
        blob = json.dumps(vm, ensure_ascii=False)
        self.assertNotIn('"score"', blob)
        self.assertNotIn("high_priority", blob)
        # 摘要的主要障碍条目已剔除 v1 分数字段
        primary = vm["summary"]["primary_obstacle"]
        self.assertEqual(primary["dim"], "A1")
        self.assertNotIn("score", primary)
        # 缺失 priority_order 时按固定 A1–A5 顺序降级
        self.assertEqual(vm["summary"]["priority_order"], ["A1", "A2"])
        # v1 缺 support_plan：A6 卡不报错
        a6 = cards_by_dim(vm)["A6"]
        self.assertEqual(a6["state_label"], "支持选择器")


class V1SequenceNoteNotRendered(unittest.TestCase):
    """审前裁定：v1 legacy 旧权重叙事（分数 ×0.7 等）不得透传进视图模型。

    v1/缺 priority_order 的 profile：摘要 sequence_note 置空；
    v1 cross_mods.action 中冻结的旧打分措辞回退为规则表现行规范文本；
    v2 的规则化 sequence_note（priority_order 存在）正常显示。
    """

    def test_v1_sequence_note_with_score_wording_is_dropped(self):
        v1 = {
            "schema": "six-am-profile-v1",
            "student": {"name": "旧版学生", "grade": "六年级"},
            "report_id": 8,
            "dimensions": {
                "A3": {"state": "At-Risk", "confidence": "Hypothesis",
                       "score": {"total": 1.4},
                       "triggered_rules": [], "cross_mods": [], "flags": [],
                       "notes": []},
                "A4": {"state": "At-Risk", "confidence": "Hypothesis",
                       "score": {"total": 0.98},
                       "triggered_rules": [],
                       # v1 旧引擎冻结的跨维修正文案，含打分叙事（真实库中亦存在）
                       "cross_mods": [{"id": "X.3",
                                       "action": "A3 优先排除：obstacle 排序中 A3 强制置顶，"
                                                 "A4 score ×0.7（其分数让位于 A3）"}],
                       "flags": [], "notes": []},
                "A5": {"state": "Secure", "score": {"total": 0.0},
                       "triggered_rules": [],
                       "cross_mods": [{"id": "X.3",
                                       "action": "A3 优先排除：obstacle 排序中 A3 强制置顶，"
                                                 "A5 score ×0.7（其分数让位于 A3）"}],
                       "flags": [], "notes": []},
                # v1 A6 支持排序条目带数值 score 字段（真实库中亦存在）
                "A6": {"state": None,
                       "support_ranking": {
                           "ranked": [{"direction": "家长陪同",
                                       "sub_mechanisms": ["陪跑"],
                                       "predicted_effectiveness": "高",
                                       "score": 3.0,
                                       "basis_rules": ["R6.8"],
                                       "cautions": ""}],
                           "avoid": [{"direction": "权威施压", "kind": "避免",
                                      "basis_rules": ["R6.6"],
                                      "reason": "高压 Accountability 避免"}],
                       }},
            },
            # v1：无 priority_order，sequence_note 为旧打分叙事
            "sequence_note": "X.3 命中：A3 强制置顶；A4/A5 分数 ×0.7（让位于 A3），"
                             "其 Hypothesis 级结论待 A3 处理后复测。",
            "obstacles": [{"dim": "A3", "score": 1.4, "high_priority": False,
                           "sequence": 1, "sequence_note": "X.3 强制置顶"}],
        }
        vm = presentation.build_view_model(v1)
        self.assertEqual(vm["summary"]["sequence_note"], "")
        blob = json.dumps({"summary": vm["summary"], "cards": vm["cards"]},
                          ensure_ascii=False)
        self.assertNotIn("分数", blob)
        self.assertNotIn("×0.7", blob)
        self.assertNotIn('"score"', blob)
        # cross_mods 保留条目与规则 ID，action 已回退为现行规范文本（无打分措辞）
        a4_mods = cards_by_dim(vm)["A4"]["audit"]["cross_mods"]
        self.assertEqual([m["id"] for m in a4_mods], ["X.3"])
        self.assertTrue(a4_mods[0]["action"])
        self.assertNotIn("score", a4_mods[0]["action"].lower())
        # A6 支持排序：数值 score 剔除，定性字段保留
        ranked = cards_by_dim(vm)["A6"]["audit"]["support_ranking"]["ranked"]
        self.assertEqual(ranked[0]["direction"], "家长陪同")
        self.assertEqual(ranked[0]["predicted_effectiveness"], "高")
        self.assertNotIn("score", ranked[0])

    def test_v2_rule_based_sequence_note_is_kept(self):
        evidence = {"ft": {"schedule": {"load_level": "高", "free_hours_per_week": 2}}}
        profile = run(overload_items(), evidence=evidence)
        # v2 引擎产出：priority_order 存在，sequence_note 为规则化说明（无分数叙事）
        self.assertTrue(profile.get("priority_order"))
        self.assertTrue(profile.get("sequence_note"))
        vm = presentation.build_view_model(profile)
        self.assertEqual(vm["summary"]["sequence_note"], profile["sequence_note"])
        self.assertIn("X.3 命中", vm["summary"]["sequence_note"])


def _docx_text(profile):
    """渲染 DOCX 并抽取全部段落与表格文本，供断言。"""
    import io
    from docx import Document

    from six_am.render_docx import build_six_am_docx

    blob = build_six_am_docx(profile)
    doc = Document(io.BytesIO(blob))
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


class DocxActionCards(unittest.TestCase):
    """Task 6：Word 档案镜像网页行动卡，不再出现 Score / ConfidenceFactor / Amplifier。"""

    FORBIDDEN = ("ObstacleScore", "ConfidenceFactor", "Amplifier", "分数 ×")

    def test_docx_contains_action_card_sections_and_no_obstacle_score(self):
        profile = engine.run(
            rebellion_items(),
            student={"name": "测试学生", "grade": "高二"},
            evidence={"ft": {"schedule": {"load_level": "高", "free_hours_per_week": 2}}},
        )
        text = _docx_text(profile)
        # 每维行动卡字段
        for label in ("当前判断", "证据状态", "Y4 事实", "初筛假设", "补充证据",
                      "下一步", "避免事项", "待确认问题"):
            self.assertIn(label, text, label)
        # 固定章节
        for section in ("一、紧凑摘要", "二、六维行动卡", "三、待确认问题与数据缺口",
                        "四、附录：规则审计"):
            self.assertIn(section, text, section)
        # 六维固定顺序：按各卡标题定位（避免章节标题 "A1–A6" 干扰）
        vm = presentation.build_view_model(profile)
        heads = ["%s %s · %s" % (c["dim"], c["title"], c["state_label"])
                 for c in vm["cards"]]
        idx = [text.index(h) for h in heads]
        self.assertEqual(idx, sorted(idx), heads)
        # 不得出现旧 Score 术语
        for bad in self.FORBIDDEN:
            self.assertNotIn(bad, text, bad)

    def test_v1_profile_docx_carries_no_score_wording(self):
        v1 = {
            "schema": "six-am-profile-v1",
            "student": {"name": "旧版学生", "grade": "高一"},
            "report_id": 7,
            "dimensions": {
                "A1": {"state": "Blocked", "subtype": "A1-S2",
                       "subtype_name": "Rebellion 逆反型（高叛逆风险）",
                       "confidence": "Hypothesis",
                       "score": {"total": 12.3, "base": 3.0},
                       "triggered_rules": [
                           {"id": "R1.3", "title": "高叛逆风险（双条件）",
                            "evidence": [{"point": "自驱力-自主性", "eval_value": "偏低"}],
                            "effect": {"state": "Blocked"}, "confidence": "Hypothesis",
                            "contributing": True}],
                       "cross_mods": [], "flags": [], "notes": []},
                "A2": {"state": "Secure", "score": {"total": 0.0},
                       "triggered_rules": [], "cross_mods": [], "flags": [], "notes": []},
                "A3": {"state": "Secure", "score": {"total": 0.0},
                       "triggered_rules": [], "cross_mods": [], "flags": [], "notes": []},
                "A4": {"state": "Secure", "score": {"total": 0.0},
                       "triggered_rules": [], "cross_mods": [], "flags": [], "notes": []},
                "A5": {"state": "Secure", "score": {"total": 0.0},
                       "triggered_rules": [], "cross_mods": [], "flags": [], "notes": []},
                "A6": {"state": None, "score": {"total": 0.0},
                       "support_ranking": {"ranked": [
                           {"direction": "家长陪同", "sub_mechanisms": ["陪跑"],
                            "predicted_effectiveness": "高", "score": 3.0,
                            "fit_label": "优先", "basis_rules": ["R6.8"],
                            "cautions": ""}]}},
            },
            "sequence_note": "X.3 命中：A4/A5 分数 ×0.7（让位于 A3）。",
            "obstacles": [{"dim": "A1", "score": 12.3, "high_priority": True,
                           "sequence": 1, "sequence_note": ""}],
        }
        text = _docx_text(v1)
        self.assertIn("当前判断", text)
        self.assertIn("下一步", text)
        for bad in self.FORBIDDEN:
            self.assertNotIn(bad, text, bad)
        self.assertNotIn('"score"', text)
        self.assertNotIn("predicted_effectiveness", text)


if __name__ == "__main__":
    unittest.main(verbosity=2)
