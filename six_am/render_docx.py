"""6 A.M. 档案 DOCX 渲染器（six-am-docx-v1）。

输入为 engine.run 产出的 profile dict（实施计划 §3.3），确定性渲染，返回字节流，
不落盘。章节标题即解析锚点，core properties 写 schema 版本便于机器解析。

展示口径与网页一致：优先消费 presentation.build_view_model(profile)，不在 Word
内自行拼接 Score 总表或数值权重（Task 6）。

章节（固定顺序）：
  1. 标题 + 学生信息表
  2. 一、紧凑摘要
  3. 二、六维行动卡（A1–A6 固定顺序）
  4. 三、待确认问题与数据缺口
  5. 四、附录：规则审计
"""

from __future__ import annotations

from typing import Any, Dict, List, Optional

from . import presentation
from . import rules as R

FONT = "微软雅黑"
DOCX_SCHEMA = "six-am-docx-v1"

_GAP_HINTS = {
    "family_complaint": ("家长诉求（R2.8）", "主诉标签（无兴趣 / 动力不足 / 其他）+ 主诉原文 + 采集来源（家长环节/顾问记录）"),
    "schedule": ("排期（R3.6）", "一周日程概要（每天课程/活动时段数）+ 自由时间估计（小时/周）+ 负荷评级（高/中/低）"),
    "subject_status": ("学科状态（R4.9）", "逐学科条目：学科名 + 学生自评水平 + 实际状态（成绩/教师反馈）+ 偏差方向（自评低于/等于/高于实际）"),
    "standard_task": ("标准化任务结果（R5.7）", "任务形式 + 能否拆出有序步骤（是/否）+ 能否承诺第一步与开始时间（是/否）+ 情绪反应（畏难/烦躁/无所谓/正常）"),
}


def build_six_am_docx(profile: Dict[str, Any]) -> bytes:
    """从 profile dict 渲染 6 A.M. 档案 DOCX，返回文件字节。"""
    import io
    from datetime import datetime

    from docx import Document
    from docx.shared import Pt, RGBColor, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    BRAND_RED = RGBColor(0xB3, 0x3A, 0x3A)
    INK = RGBColor(0x14, 0x14, 0x14)
    GRAY = RGBColor(0x8A, 0x8A, 0x8A)
    MUTED = RGBColor(0x64, 0x74, 0x8B)

    def _set_eastasia(target) -> None:
        rPr = target._element.get_or_add_rPr()
        rFonts = rPr.find(qn("w:rFonts"))
        if rFonts is None:
            rFonts = OxmlElement("w:rFonts")
            rPr.append(rFonts)
        rFonts.set(qn("w:eastAsia"), FONT)

    def _run(p, text: str, size: float = 11, color=INK, bold: bool = False):
        r = p.add_run(str(text))
        r.font.name = FONT
        r.font.size = Pt(size)
        r.font.color.rgb = color
        r.bold = bold
        _set_eastasia(r)
        return r

    def _body(text: str, size: float = 11, color=INK, bold: bool = False,
              space_before: float = 2, space_after: float = 2,
              line: Optional[float] = 20, align=None):
        p = doc.add_paragraph()
        if align is not None:
            p.alignment = align
        pf = p.paragraph_format
        pf.space_before = Pt(space_before)
        pf.space_after = Pt(space_after)
        if line:
            pf.line_spacing = Pt(line)
        _run(p, text, size=size, color=color, bold=bold)
        return p

    def _h1(text: str):
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_before = Pt(14)
        pf.space_after = Pt(6)
        _run(p, text, size=15, color=BRAND_RED, bold=True)
        return p

    def _h3(text: str):
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_before = Pt(8)
        pf.space_after = Pt(3)
        _run(p, text, size=12, color=INK, bold=True)
        return p

    def _field(label: str, value: Any, color=INK):
        """字段行：加粗小标签 + 正文值，无值时以 "—" 占位。"""
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_before = Pt(2)
        pf.space_after = Pt(1)
        pf.line_spacing = Pt(18)
        _run(p, label + "：", size=10.5, color=MUTED, bold=True)
        text = str(value).strip() if value not in (None, "") else "—"
        _run(p, text or "—", size=10.5, color=color)
        return p

    def _field_list(label: str, items: Any, color=INK):
        """列表字段行：加粗小标签 + "· 条目"；空列表以 "—" 占位。"""
        vals = [str(x).strip() for x in (items or []) if str(x).strip()]
        if not vals:
            return _field(label, "—", color=color)
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_before = Pt(2)
        pf.space_after = Pt(1)
        pf.line_spacing = Pt(18)
        _run(p, label + "：", size=10.5, color=MUTED, bold=True)
        for i, val in enumerate(vals):
            if i:
                _run(p, "　", size=10.5)
            _run(p, "· " + val, size=10.5, color=color)
        return p

    def _clear_cell(cell):
        p = cell.paragraphs[0]
        for r in list(p.runs):
            r._element.getparent().remove(r._element)
        return p

    def _cell_text(cell, text: str, size: float = 10, color=INK, bold: bool = False, align=None):
        p = _clear_cell(cell)
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(1)
        if align is not None:
            p.alignment = align
        _run(p, text, size=size, color=color, bold=bold)
        return p

    def _shade(cell, fill: str) -> None:
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), fill)
        tcPr.append(shd)

    def _table_grid(table, color: str = "D9D9D9") -> None:
        tblPr = table._tbl.tblPr
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            el = OxmlElement(f"w:{edge}")
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), "4")
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), color)
            borders.append(el)
        tblPr.append(borders)

    def _table(rows: List[List[str]], widths: Optional[List[float]] = None,
               header: bool = True) -> None:
        if not rows:
            return
        t = doc.add_table(rows=len(rows), cols=len(rows[0]))
        _table_grid(t)
        for ri, row in enumerate(rows):
            for ci, val in enumerate(row):
                cell = t.cell(ri, ci)
                is_head = header and ri == 0
                _cell_text(cell, val, bold=is_head,
                           color=RGBColor(0xFF, 0xFF, 0xFF) if is_head else INK)
                if is_head:
                    _shade(cell, "B33A3A")
                if widths and ci < len(widths):
                    cell.width = Cm(widths[ci])
        return t

    # ── 工具：规则审计证据文本化（不含数值权重）──────────────
    def _audit_ev(e: Any) -> str:
        if not isinstance(e, dict):
            return str(e)
        if e.get("point"):
            s = "%s=%s" % (e.get("point"), e.get("eval_value"))
            if e.get("source"):
                s += "（%s）" % e["source"]
            return s
        if e.get("order_value") is not None:
            return "%s 第 %s 位/前 %s" % (e.get("order_value"), e.get("pos"), e.get("top_n"))
        if e.get("rule"):
            return "依赖 %s 命中" % e["rule"]
        if e.get("evidence"):
            return "证据 %s" % e["evidence"]
        if e.get("predicate"):
            return "谓词 %s" % e["predicate"]
        if e.get("dim"):
            return "%s state=%s" % (e.get("dim"), e.get("state"))
        return ""

    def _dim_title(dim: str) -> str:
        base = R.DIM_NAMES.get(dim, dim)
        if dim == "A6":
            return f"{dim} {base.split(' ', 1)[-1] if ' ' in base else base}"
        return f"{dim} {base}"

    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(2.2)
        section.bottom_margin = Cm(2.2)
        section.left_margin = Cm(2.2)
        section.right_margin = Cm(2.2)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = INK
    _set_eastasia(normal)

    student = profile.get("student") or {}
    rv = profile.get("rule_version") or {}
    gaps = profile.get("gaps") or {}
    answers = profile.get("answers") or []

    # 网页与 Word 唯一展示入口：视图模型（容错 v1、屏蔽旧 Score 文案）
    vm = presentation.build_view_model(profile)
    summary = vm.get("summary") or {}
    cards = vm.get("cards") or []

    # ── 1. 标题 + 学生信息表 ─────────────────────────────────
    _body("凭远教育 · 6 A.M. 六维启动档案", size=16, color=BRAND_RED, bold=True,
          space_before=0, space_after=4, line=None)
    _body(student.get("name") or "—", size=20, color=INK, bold=True,
          space_before=0, space_after=10, line=None)
    _table([
        ["学生", student.get("name") or "—", "年级", student.get("grade") or "—"],
        ["报告编号", str(profile.get("report_id") or "—"), "生成时间", profile.get("generated_at") or "—"],
        ["规则版本", (str(rv.get("label")) if rv.get("label") else "default")
         + (f"（id={rv.get('id')}）" if rv.get("id") else ""),
         "档案 Schema", profile.get("schema") or "six-am-profile-v1"],
    ], widths=[2.6, 6.4, 2.6, 6.4], header=False)

    # ── 2. 一、紧凑摘要 ──────────────────────────────────────
    _h1("一、紧凑摘要")
    primary = summary.get("primary_obstacle") or {}
    p_dim = str(primary.get("dim") or "")
    p_state = str(primary.get("state") or "")
    if p_dim:
        head = _dim_title(p_dim)
        if p_state:
            head += " · " + presentation.STATE_LABELS.get(p_state, p_state)
        if primary.get("priority_reason"):
            head += "（%s）" % primary["priority_reason"]
        _body("主要障碍：" + head, bold=True)
    else:
        _body("主要障碍：暂无（各维均在 Watch 以下）。", bold=True)
    order = summary.get("priority_order") or []
    if order:
        _body("处理顺序：" + " → ".join(_dim_title(str(d)) for d in order))
    deferred_list = summary.get("deferred") or []
    if deferred_list:
        _body("本轮暂缓（长期目标）：" + "；".join(
            "%s（%s）" % (_dim_title(str(d.get("dim"))), d.get("reason") or "")
            for d in deferred_list if isinstance(d, dict)))
    refuted = summary.get("refuted") or []
    _body("待确认问题：%s 条；被排除规则：%s 条"
          % (summary.get("open_question_count") or 0, len(refuted)))
    if summary.get("sequence_note"):
        _body(summary["sequence_note"], size=10, color=MUTED, space_before=4)

    # ── 3. 二、六维行动卡（A1–A6 固定顺序）──────────────────
    _h1("二、六维行动卡（A1–A6）")
    for card in cards:
        head = "%s %s · %s" % (card.get("dim"), card.get("title"), card.get("state_label"))
        if card.get("manually_edited"):
            head += "（人工修订）"
        _h3(head)
        _field("当前判断", card.get("headline"))
        _field("证据状态", card.get("evidence_label"))
        _field_list("Y4 事实", card.get("y4_facts"))
        _field_list("初筛假设", card.get("hypotheses"))
        _field_list("补充证据", card.get("confirmed_evidence"))
        if card.get("priority_label"):
            if card.get("deferred"):
                prio = "长期目标"
            elif card.get("priority_rank") is not None:
                prio = "优先处理 · 第 %s 位" % card["priority_rank"]
            else:
                prio = card["priority_label"]
            if card.get("priority_reason") and card.get("priority_reason") != prio:
                prio += "（%s）" % card["priority_reason"]
            _field("处理优先级", prio)
        _field_list("下一步", card.get("next_actions"))
        _field_list("避免事项", card.get("avoid_actions"))
        _field_list("待确认问题",
                    [q.get("ask") for q in card.get("questions") or [] if q.get("ask")])

    # ── 4. 三、待确认问题与数据缺口 ──────────────────────────
    _h1("三、待确认问题与数据缺口")
    questions = profile.get("questions") or []
    if questions:
        rows = [["#", "怎么问（人话版）", "验证目标", "关联规则"]]
        for q in questions:
            rows.append([
                str(q.get("priority") or ""),
                str(q.get("ask") or ""),
                str(q.get("id") or ""),
                "、".join(q.get("basis_rules") or q.get("verify") or []) or "—",
            ])
        _table(rows, widths=[1.2, 9.6, 2.6, 5.6])
    else:
        _body("当前无需要追问的维度（各维均低于 Watch）。")
    if answers:
        _body("已录入追问答案：", bold=True, space_before=6)
        for a in answers:
            mark = {"confirm": "证实", "partial": "部分证实", "deny": "否定"}.get(
                str(a.get("result")), str(a.get("result") or ""))
            _body(f"· {a.get('q')} [{mark}] {a.get('text') or ''}", size=10)
    missing = [k for k in _GAP_HINTS if not gaps.get(k)]
    if missing:
        _body("数据缺口与补录指引：", bold=True, space_before=6)
        for k in missing:
            label, hint = _GAP_HINTS[k]
            _body(f"· {label}：尚未采集。补录：{hint}", size=10)
    else:
        _body("四类结构化数据（家长诉求 / 排期 / 学科状态 / 标准化任务）均已采集。", size=10)
    collected = [k for k in _GAP_HINTS if gaps.get(k)]
    if collected:
        _body("已采集字段：", color=MUTED, space_before=4)
        for k in collected:
            _body(f"· {_GAP_HINTS[k][0]}：{gaps.get(k)}", size=9.5, color=MUTED)

    # ── 5. 四、附录：规则审计 ────────────────────────────────
    _h1("四、附录：规则审计")
    for card in cards:
        audit = card.get("audit") or {}
        traces = audit.get("triggered_rules") or []
        mods = audit.get("cross_mods") or []
        flags = audit.get("flags") or []
        notes = audit.get("notes") or []
        refuted_rules = audit.get("refuted_rules") or []
        if not (traces or mods or flags or notes or refuted_rules):
            continue
        _h3("%s %s" % (card.get("dim"), card.get("title")))
        for t in traces:
            if not isinstance(t, dict):
                continue
            line = "%s %s" % (t.get("id"), t.get("title") or "")
            marks = []
            if t.get("contributing") is False:
                marks.append("仅记录不产出")
            if t.get("refuted"):
                marks.append("refuted")
            if marks:
                line += "（%s）" % "、".join(marks)
            evs = [s for s in (_audit_ev(e) for e in (t.get("evidence") or [])) if s]
            if evs:
                line += "：" + "；".join(evs)
            _body("· " + line, size=9.5)
        for m in mods:
            _body("· 跨维修正 %s：%s" % (m.get("id"), m.get("action") or ""), size=9.5, color=MUTED)
        for fl in flags:
            _body("· 标记：%s" % fl, size=9.5, color=MUTED)
        for n in notes:
            _body("· 注：%s" % n, size=9.5, color=MUTED)
        for r in refuted_rules:
            _body("· %s %s %s（由追问 %s 否定）"
                  % (r.get("rule"), r.get("title"), r.get("label"), r.get("refuted_by")),
                  size=9.5, color=MUTED)

    # ── 页脚 + core properties ───────────────────────────────
    doc.add_paragraph()
    fp = doc.add_paragraph()
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _run(fp, "— 凭远教育 · 6 A.M. 六维启动模型 —", size=9, color=GRAY)

    cp = doc.core_properties
    cp.title = f"6AM 六维启动档案 - {student.get('name') or ''}"
    cp.subject = DOCX_SCHEMA
    cp.keywords = f"{DOCX_SCHEMA};{profile.get('schema') or 'six-am-profile-v1'}"
    cp.comments = "6 A.M. 六维启动模型档案（v0.2），章节标题为解析锚点。"

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()